"""Listing photos: five 2500 x 2000 images per product, drawn from the real files.

1  the cover: the product name and its pages fanned out
2  what's inside: every page as a thumbnail
3  a close look: one page up close, with what makes it useful
4  how you use it: on an iPad, a laptop, a phone or on paper
5  made in a real studio: the maker and the promises
Also the shop icon (500 x 500) and banner (3360 x 840).
"""
import base64, html, os, re, shutil

import pymupdf

from brand import TOKENS, font_css

HERE = os.path.dirname(os.path.abspath(__file__))
ART = os.path.join(HERE, "art")
W, H = 1250, 1000                      # CSS px; shot at 2x


def b64(path, mime=None):
    mime = mime or ("image/png" if path.endswith(".png") else "image/jpeg")
    return "data:%s;base64,%s" % (mime, base64.b64encode(open(path, "rb").read()).decode())


def esc(s):
    return html.escape(str(s))


def rich(s):
    s = esc(s)
    return re.sub(r"_(.+?)_", r"<i>\1</i>", s)


CSS = """
*{box-sizing:border-box}
html,body{margin:0;background:#ddd}
.art{width:1250px;height:1000px;position:relative;overflow:hidden;font-family:var(--sans);color:var(--ink);
  background:radial-gradient(120% 90% at 85% 15%,#fbf7f7 0%,#f3e9eb 55%,#ecdfe2 100%)}
.art+.art{margin-top:10px}
.mark{position:absolute;right:-40px;bottom:-60px;width:420px;opacity:.10}
.brandline{position:absolute;left:70px;top:58px;display:flex;align-items:center;gap:16px;font-size:13px;
  letter-spacing:.32em;text-transform:uppercase;color:var(--rose);font-weight:500}
.brandline span{width:44px;height:1px;background:var(--rose);opacity:.5}
.brandline em{font-style:normal;color:var(--muted);letter-spacing:.2em;font-weight:400}
h1{font-family:var(--serif);font-weight:500;font-size:70px;line-height:.98;letter-spacing:-.01em;margin:0;text-wrap:balance}
h1 i{color:var(--rose)}
.sub{font-size:21px;line-height:1.45;color:var(--muted);font-weight:300;margin:22px 0 0}
.chips{display:flex;flex-wrap:wrap;gap:10px;margin-top:34px}
.chip{border:1.2px solid rgba(168,107,126,.55);color:var(--ink);border-radius:999px;padding:9px 18px 8px;
  font-size:13px;letter-spacing:.14em;text-transform:uppercase;font-weight:400;background:rgba(255,255,255,.55)}
.sheet{position:absolute;background:#fff;box-shadow:0 34px 60px -24px rgba(90,45,62,.42),0 4px 12px rgba(60,30,40,.10)}
.sheet img{display:block;width:100%;height:100%}
.copy{position:absolute;left:70px;top:150px;width:470px}
.price{position:absolute;left:70px;bottom:64px;font-size:15px;letter-spacing:.18em;text-transform:uppercase;color:var(--muted)}
.price b{font-family:var(--serif);font-size:30px;letter-spacing:0;color:var(--ink);font-weight:500;text-transform:none;margin-right:10px}
.top{position:absolute;left:70px;right:70px;top:120px;text-align:center}
.top h1{font-size:62px}
.grid{position:absolute;left:70px;right:70px;top:290px;bottom:60px;display:grid;gap:26px 22px;align-content:start}
.tile{display:flex;flex-direction:column;align-items:center;gap:12px;min-width:0}
.tile .sheet{position:relative;width:100%}
.tile p{margin:0;font-size:14.5px;letter-spacing:.06em;text-align:center;color:var(--ink);line-height:1.3}
.feat{position:absolute;right:60px;top:150px;width:430px;display:flex;flex-direction:column;gap:40px}
.feat div{display:grid;grid-template-columns:44px 1fr;gap:16px;align-items:start}
.feat b{width:44px;height:44px;border-radius:50%;background:var(--rose);color:#fff;display:flex;align-items:center;
  justify-content:center;font-family:var(--serif);font-size:24px;font-weight:500}
.feat h3{margin:2px 0 8px;font-family:var(--serif);font-size:34px;font-weight:600;line-height:1.05}
.feat p{margin:0;font-size:20px;line-height:1.45;color:var(--muted);font-weight:300}
.crop{position:absolute;left:70px;top:70px;width:640px;height:860px;overflow:hidden;background:#fff;border-radius:6px;
  box-shadow:0 34px 60px -24px rgba(90,45,62,.42),0 4px 12px rgba(60,30,40,.10)}
.crop img{position:absolute;display:block}
.pin{position:absolute;width:40px;height:40px;border-radius:50%;background:var(--rose);color:#fff;display:flex;
  align-items:center;justify-content:center;font-family:var(--serif);font-size:22px;font-weight:500;
  box-shadow:0 0 0 7px rgba(168,107,126,.22)}
.ipad{position:absolute;background:#1f1c1d;border-radius:44px;padding:26px;box-shadow:0 40px 70px -28px rgba(40,20,28,.55)}
.ipad .scr{width:100%;height:100%;border-radius:16px;overflow:hidden;background:#fff;position:relative}
.ipad .scr img{width:100%;display:block}
.phone{position:absolute;background:#1f1c1d;border-radius:46px;padding:12px;box-shadow:0 40px 70px -28px rgba(40,20,28,.55)}
.phone .scr{width:100%;height:100%;border-radius:36px;overflow:hidden;background:#fff}
.phone .scr img{width:100%;height:100%;display:block;object-fit:cover}
.laptop{position:absolute;left:330px;top:250px;width:860px}
.laptop .lid{background:#1f1c1d;border-radius:22px 22px 0 0;padding:18px 18px 22px}
.laptop .scr{background:#fff;height:500px;overflow:hidden;border-radius:4px}
.laptop .base{height:22px;margin:0 -46px;background:linear-gradient(#d9d2d4,#bfb5b8);border-radius:0 0 18px 18px}
.laptop .base::after{content:"";display:block;width:140px;height:8px;margin:0 auto;background:#a79ca0;border-radius:0 0 8px 8px}
.xl{font-family:Arial,Helvetica,sans-serif;font-size:11px;color:#222;height:100%;display:flex;flex-direction:column}
.xl .bar{height:30px;background:#f1eff0;border-bottom:1px solid #d9d4d6;display:flex;align-items:center;gap:7px;padding:0 12px;font-size:11.5px;color:#555}
.xl .bar i{width:10px;height:10px;border-radius:50%;background:#d9cfd2;display:inline-block}
.xl .bar span{margin-left:12px}
.xl .fx{height:24px;border-bottom:1px solid #e3dfe0;display:flex;align-items:center;padding:0 10px;color:#777;font-size:11px;gap:14px}
.xl .fx b{color:#999;font-weight:400;font-style:italic}
.xl .body{flex:1;overflow:hidden;position:relative}
.xl table{border-collapse:collapse;table-layout:fixed}
.xl td,.xl th{border:1px solid #e6e2e3;height:20px;padding:0 5px;white-space:nowrap;overflow:hidden;text-overflow:clip}
.xl th{background:#f6f4f5;color:#888;font-weight:400;font-size:10px;text-align:center}
.xl .tabs{height:28px;border-top:1px solid #d9d4d6;background:#f6f4f5;display:flex;align-items:stretch;font-size:11px}
.xl .tabs span{padding:7px 14px;color:#666;border-right:1px solid #e1dcde}
.xl .tabs span.on{background:#fff;color:#a86b7e;font-weight:bold;border-bottom:2px solid #a86b7e}
.maker{position:absolute;left:0;top:0;right:0;bottom:0;display:grid;grid-template-columns:470px 1fr}
.maker .pic{position:relative;background:linear-gradient(160deg,#ead9de,#f6eef0)}
.maker .pic img{position:absolute;left:50%;top:50%;width:290px;transform:translate(-50%,-46%);opacity:.95}
.maker .txt{padding:150px 80px 0 70px}
.maker h1{font-size:64px}
.promise{margin-top:40px;display:grid;gap:22px}
.promise div{display:grid;grid-template-columns:34px 1fr;gap:14px;font-size:19px;line-height:1.4;font-weight:300}
.promise div::before{content:"";width:22px;height:22px;margin-top:3px;border-radius:50%;border:1.5px solid var(--rose);
  background:radial-gradient(circle,var(--rose) 0 4px,transparent 5px)}
"""


def page_png(pdf, idx, out, dpi=150):
    if not os.path.exists(out):
        d = pymupdf.open(pdf)
        d[idx].get_pixmap(dpi=dpi, alpha=False).save(out)
        d.close()
    return out


def fmt(v, nf):
    if v is None:
        return ""
    if isinstance(v, str):
        return v
    if hasattr(v, "strftime"):
        return v.strftime("%m/%d/%Y")
    if isinstance(v, bool):
        return str(v)
    zero_dash = '"-"' in nf
    if zero_dash and abs(v) < 1e-9:
        return "-"
    neg = v < 0
    a = abs(v)
    if "%" in nf:
        dec = 1 if "0.0%" in nf else 0
        s = ("{:,.%df}%%" % dec).format(a * 100)
    elif "$" in nf:
        dec = 2 if ".00" in nf else 0
        s = "$" + ("{:,.%df}" % dec).format(a)
    elif "#,##0.0" in nf:
        s = "{:,.1f}".format(a)
    elif "#,##0" in nf or nf == "0":
        s = "{:,.0f}".format(a)
    else:
        s = ("%g" % a)
    if neg:
        s = "(%s)" % s if "(" in nf else "-" + s
    return s


def sheet_html(xlsx, name, rows, cols, tabs):
    """A spreadsheet window showing one sheet as it looks in Excel."""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    wf = load_workbook(xlsx)
    wv = load_workbook(xlsx, data_only=True)
    ws, vs = wf[name], wv[name]
    cf = []
    for rng in ws.conditional_formatting:
        for cr in rng.sqref.ranges:
            cf.append(cr.bounds)                       # (min_col, min_row, max_col, max_row)
    widths = []
    for c in range(1, cols + 1):
        dim = ws.column_dimensions[get_column_letter(c)]
        widths.append(int((dim.width or 9) * 6.4))
    out = ['<table><tr><th style="width:30px"></th>%s</tr>' % "".join(
        '<th style="width:%dpx">%s</th>' % (w, get_column_letter(i + 1)) for i, w in enumerate(widths))]
    for r in range(1, rows + 1):
        h = ws.row_dimensions[r].height
        cells = []
        for c in range(1, cols + 1):
            cell = ws.cell(row=r, column=c)
            v = vs.cell(row=r, column=c).value
            st = []
            f = cell.font
            if f is not None:
                if f.color is not None and isinstance(f.color.rgb, str):
                    st.append("color:#%s" % f.color.rgb[-6:])
                if f.b:
                    st.append("font-weight:bold")
                if f.i:
                    st.append("font-style:italic")
                if f.sz and f.sz != 10:
                    st.append("font-size:%.1fpx" % (f.sz * 1.1))
            if cell.fill is not None and cell.fill.fill_type == "solid" and isinstance(cell.fill.fgColor.rgb, str):
                st.append("background:#%s" % cell.fill.fgColor.rgb[-6:])
            for (c1, r1, c2, r2) in cf:
                if c1 <= c <= c2 and r1 <= r <= r2 and isinstance(v, (int, float)):
                    st.append("background:%s" % ("#f8dcdc" if v < 0 else "#e3f1e3"))
            text = fmt(v, cell.number_format or "")
            if isinstance(v, (int, float)) or (cell.alignment is not None and cell.alignment.horizontal == "right"):
                st.append("text-align:right")
            cells.append('<td style="%s">%s</td>' % (";".join(st), esc(text)))
        out.append('<tr style="height:%dpx"><th>%d</th>%s</tr>' % (int((h or 15) * 1.33), r, "".join(cells)))
    out.append("</table>")
    tabbar = "".join('<span class="%s">%s</span>' % ("on" if t == name else "", esc(t)) for t in tabs)
    return ('<div class="xl"><div class="bar"><i></i><i></i><i></i><span>%s</span></div>'
            '<div class="fx"><b>fx</b></div><div class="body">%s</div><div class="tabs">%s</div></div>'
            % (esc(os.path.basename(xlsx)), "".join(out), tabbar))


# ─────────────────────────── per-product picture plans ───────────────────────────

def plans(dist, tmp):
    f = lambda slug, name: os.path.join(dist, slug, "files", name)
    P = {}
    s1 = "01-intake-consent-kit"
    pdf1 = f(s1, "Intake & Consent Kit - Fillable - US Letter.pdf")
    P["01"] = dict(
        pdf=pdf1, fan=[0, 1, 3], headline="Client Intake _& Consent_ Kit",
        sub="Eight forms every new client signs, written by a licensed esthetician.",
        chips=["Fillable PDF", "Word", "Letter + A4"],
        inside=[(pdf1, i, n) for i, n in enumerate(["Client Intake", "Health History", "Facial Consent",
                                                    "Chemical Peel Consent", "Dermaplaning Consent", "Waxing Consent",
                                                    "Photo Release", "Studio Policies"])],
        detail=(pdf1, 3, (0.0, 0.03, 0.94), ["Before your peel", "What to expect", "Patch test"],
                [("Initial each statement", "Clients initial every risk and instruction, so nothing is missed."),
                 ("Before and after", "What to avoid before a peel, and exactly what to expect after."),
                 ("Patch test and notes", "Record the patch test, the peel, layers and how the skin responded.")]),
        use="ipad", use_title="Type on an iPad, _or print it_",
        use_sub="Every line and box is a fillable field. Clients fill it in, sign with a finger or Apple Pencil, and you save it to their file.")
    s2 = "02-aftercare-cards"
    pdf2 = f(s2, "Aftercare Cards - 5x7 Print Sheets - US Letter.pdf")
    P["02"] = dict(
        pdf=pdf2, landscape=True, fan=[0, 1, 2], headline="Aftercare _Cards_",
        sub="Eight cards that answer the questions clients have the night after a treatment.",
        chips=["5 x 7 print", "Phone versions", "Word"],
        inside="phone", detail=None, use="phones",
        use_title="Hand it over, _or text it_",
        use_sub="Every card also comes as a phone-sized image, ready to text before your client reaches the car.")
    s3 = "03-skin-analysis-face-map"
    pdf3 = f(s3, "Skin Analysis & Treatment Record - Fillable - US Letter.pdf")
    P["03"] = dict(
        pdf=pdf3, fan=[1, 0, 2], headline="Skin Analysis _& Face Map_",
        sub="A professional consultation on one page, plus a visit log and a progress map.",
        chips=["Fillable PDF", "Word", "Letter + A4"],
        inside=[(pdf3, 0, "Skin Analysis"), (pdf3, 1, "Treatment Record"), (pdf3, 2, "Progress Map")],
        detail=(pdf3, 0, (0.0, 0.03, 0.94), ["Zones:", "Skin type", "In the studio"],
                [("A clean face map", "Nine numbered zones and a marking key for comedones, milia, pigment and more."),
                 ("Everything at a glance", "Skin type, Fitzpatrick, Glogau, hydration, oil, sensitivity and pores."),
                 ("The plan", "In-studio treatments, home care and the next visit, agreed before they leave.")]),
        use="ipad", use_title="Map it with _Apple Pencil_",
        use_sub="Open it in Markup or Acrobat and draw straight on the face map. Or print a stack for the treatment room.")
    s4 = "04-service-menu-templates"
    pdf4 = f(s4, "Service Menu - Sample - US Letter.pdf")
    lib4 = f(s4, "Description Library - 40 Services.pdf")
    P["04"] = dict(
        pdf=pdf4, fan=[1, 0, 2], headline="Service Menu _& Price List_",
        sub="An elegant menu you can make yours in ten minutes, and 40 descriptions already written.",
        chips=["Editable Word", "Google Docs", "40 descriptions"],
        inside=[(pdf4, 0, "Service menu, page 1"), (pdf4, 1, "Service menu, page 2"), (pdf4, 2, "Price list"),
                (lib4, 0, "Description library"), (lib4, 1, "Facials and more"), (lib4, 2, "Waxing, lash and brow")],
        detail=(pdf4, 0, (0.0, 0.03, 0.94), ["Your Studio Name", "Signature Facial", "Peels and advanced"],
                [("Your name, front and center", "Change the studio name, city and contact line in seconds."),
                 ("Descriptions that sell", "Clear, benefit-first copy with time and price side by side."),
                 ("Sections that flow", "Facials, peels, add-ons, waxing, lash and brow, then the membership.")]),
        use="print", use_title="Print it, _post it, send it_",
        use_sub="Export a PDF for your booking page and link in bio, or print it for the treatment room.")
    s5 = "05-membership-pricing-calculator"
    x5 = f(s5, "Membership Pricing Calculator.xlsx")
    P["05"] = dict(
        xlsx=x5, sheets=[("Calculator", 48, 5), ("Your First Year", 29, 11), ("Price Test", 17, 8)],
        tabs=["Start Here", "Calculator", "Your First Year", "Price Test"],
        headline="Membership _Pricing_ Calculator",
        sub="What each member is worth, how many cover your bills, and how many pay you.",
        chips=["Excel", "Google Sheets", "Numbers"], use="laptop",
        detail_feats=[("Your costs and tiers", "Up to three tiers with price, product cost, perks and member mix."),
                      ("The answers", "Break-even members, members for your pay goal, and hours they fill."),
                      ("A 12-month plan", "Sign-ups and cancellations month by month, and when you break even.")])
    s6 = "06-bookkeeping-profit-tracker"
    x6 = f(s6, "Bookkeeping & Profit Tracker.xlsx")
    P["06"] = dict(
        xlsx=x6, sheets=[("Month by Month", 38, 14), ("Income", 24, 10), ("The Year at a Glance", 16, 3)],
        tabs=["Start Here", "Settings", "Income", "Expenses", "Products", "Month by Month", "The Year at a Glance"],
        headline="Bookkeeping _& Profit_ Tracker",
        sub="Log a sale in seconds. See what you actually keep, every month.",
        chips=["Excel", "Google Sheets", "Numbers"], use="laptop",
        detail_feats=[("Drop-downs", "Categories and payment methods pick from a list. Card fees fill themselves in."),
                      ("Month by month", "Income, expenses, profit and your tax set-aside, all year."),
                      ("Your shelf", "Retail margins, stock value and reorder alerts.")])
    s7 = "07-rebooking-membership-scripts"
    pdf7 = f(s7, "Rebooking & Membership Scripts - US Letter.pdf")
    P["07"] = dict(
        pdf=pdf7, fan=[1, 0, 2], headline="Rebooking _& Membership_ Scripts",
        sub="About 45 scripts for the room, your texts and your DMs. Never pushy.",
        chips=["PDF", "Editable Word", "Letter + A4"],
        inside=[(pdf7, i, n) for i, n in enumerate(["In the room", "The membership", "Texts that fill your book",
                                                    "Protecting your time", "DMs and questions"])],
        detail=(pdf7, 1, (0.0, 0.03, 0.94), ["Introducing it", "The questions everyone asks", "The follow-up text"],
                [("Word for word", "Exactly what to say, with [brackets] for your prices and policies."),
                 ("The hard questions", "“Can I cancel?” “It's a lot every month.” Answered kindly."),
                 ("The follow-up", "The text to send the next day, so the yes doesn't slip away.")]),
        use="phones_text", use_title="Paste them into _your booking app_",
        use_sub="Confirmations, reminders and follow-ups, ready for automatic messages or text replacements.")
    s8 = "08-open-your-studio-planner"
    pdf8 = f(s8, "Open Your Studio 90-Day Planner - US Letter.pdf")
    P["08"] = dict(
        pdf=pdf8, fan=[3, 0, 6], headline="Open Your _Studio_",
        sub="A 90-day planner: the checklists, the numbers and one clear focus a week.",
        chips=["23 pages", "Fillable PDF", "Letter + A4"],
        inside=[(pdf8, 0, "Start"), (pdf8, 1, "Before you open"), (pdf8, 2, "The room"), (pdf8, 3, "Your prices"),
                (pdf8, 4, "First 25 clients"), (pdf8, 5, "30 post ideas"), (pdf8, 6, "13 weekly pages"),
                (pdf8, 10, "Monthly reviews")],
        detail=(pdf8, 6, (0.0, 0.03, 0.94), ["This week:", "Bookings", "Posts this week"],
                [("One focus a week", "Paperwork, soft launch, reviews, membership, partners: in the right order."),
                 ("Bookings and money", "Goals, new clients, rebooks, money in and out."),
                 ("Posts and wins", "Plan the week's content and keep track of what worked.")]),
        use="ipad", use_title="Plan on paper _or on an iPad_",
        use_sub="Print and bind it, or open it in GoodNotes or Acrobat and type straight on the pages.")
    s9 = "09-skin-journal-routine-planner"
    pdf9 = f(s9, "Skin Journal - US Letter.pdf")
    P["09"] = dict(
        pdf=pdf9, fan=[1, 0, 4], headline="My Skin _Journal_",
        sub="Thirty days to know your skin: what you use, how it responds, and what works.",
        chips=["14 pages", "Printable", "Fillable"],
        inside=[(pdf9, 0, "Start"), (pdf9, 1, "My routine"), (pdf9, 2, "My products"), (pdf9, 3, "Trying something new"),
                (pdf9, 4, "30-day diary"), (pdf9, 6, "Weekly check-ins"), (pdf9, 10, "Progress photos"),
                (pdf9, 11, "Habit tracker")],
        detail=(pdf9, 1, (0.0, 0.03, 0.94), ["Morning", "Evening", "Weekly"],
                [("Your morning, in order", "Cleanse to sunscreen, thinnest to thickest, with room for notes."),
                 ("Your evening", "Makeup off, treatments and serums, and what to keep on separate nights."),
                 ("Your week", "Exfoliation and masks on the right days, so nothing clashes.")]),
        use="ipad", use_title="By the sink, _or on a tablet_",
        use_sub="Print it and keep it by the mirror, or fill it in on an iPad. Bring it to your next facial.")
    P["10"] = dict(bundle=True, headline="The Esthetician _Business Kit_",
                   sub="Eight products and a bonus: the forms, the cards, the menu, the numbers and the plan.",
                   chips=["9 products", "Save 49%", "Instant download"])
    return P


def fan(pngs, landscape=False):
    out = []
    if landscape:
        spots = [(600, 200, 540, 417, -7), (680, 350, 540, 417, 4), (630, 510, 540, 417, -2)]
    else:
        spots = [(560, 170, 420, 543, -8), (800, 220, 420, 543, 7), (660, 370, 430, 556, -1)]
    for (x, y, w, h, rot), p in zip(spots, pngs):
        out.append('<div class="sheet" style="left:%dpx;top:%dpx;width:%dpx;height:%dpx;transform:rotate(%sdeg)">'
                   '<img src="%s"></div>' % (x, y, w, h, rot, b64(p)))
    return "".join(out)


def brandline(extra="Made by a licensed esthetician"):
    return '<div class="brandline">Lumevina Studio<span></span><em>%s</em></div>' % esc(extra)


def chips(items):
    return '<div class="chips">%s</div>' % "".join('<span class="chip">%s</span>' % esc(c) for c in items)


def art(inner):
    return '<div class="art"><img class="mark" src="%s">%s</div>' % (b64(os.path.join(ART, "mark-rose.png")), inner)


def doc(arts):
    return ('<!doctype html><html><head><meta charset="utf-8"><style>%s%s%s</style></head><body>%s</body></html>'
            % (font_css(), TOKENS, CSS, "".join(arts)))


def cover(p, pngs, price):
    inner = (brandline() + '<div class="copy"><h1>%s</h1><p class="sub">%s</p>%s</div>' % (
        rich(p["headline"]), esc(p["sub"]), chips(p["chips"])) + fan(pngs, p.get("landscape")) +
        '<div class="price"><b>Instant download</b>edit, print, use today</div>')
    return art(inner)


def inside(p, tiles, cols=None, title="What&rsquo;s <i>inside</i>"):
    n = len(tiles)
    tall = any(ar == "9/16" for _, _, ar in tiles)
    if tall:
        cols, w = 4, 150
    elif n <= 3:
        cols, w = n, 300
    elif n <= 6:
        cols, w = 3, 225
    elif n <= 8:
        cols, w = 4, 200
    else:
        cols, w = 5, 180
    t = "".join('<div class="tile"><div class="sheet" style="aspect-ratio:%s"><img src="%s"></div><p>%s</p></div>'
                % (ar, b64(png), esc(label)) for png, label, ar in tiles)
    return art(brandline("What's inside") + '<div class="top"><h1>%s</h1></div>'
               '<div class="grid" style="grid-template-columns:repeat(%d,%dpx);justify-content:center;top:%dpx">%s</div>'
               % (title, cols, w, 250 if n > 3 else 300, t))


def find_pins(pdf, idx, strings):
    """Where each heading sits on the page, as fractions of the page (x of its left edge, y of its middle)."""
    d = pymupdf.open(pdf)
    pg = d[idx]
    w, h = pg.rect.width, pg.rect.height
    out = []
    lines = []
    for b in pg.get_text("dict")["blocks"]:
        for ln in b.get("lines", []):
            text = "".join(sp["text"] for sp in ln["spans"])
            lines.append((re.sub(r"\s+", "", text).upper(), pymupdf.Rect(ln["bbox"])))
    for s_ in strings:
        key = re.sub(r"\s+", "", s_).upper()
        hits = sorted((r for t, r in lines if key in t), key=lambda r: (r.x0 > w / 2, r.y0))
        if not hits:
            raise SystemExit("mockup pin text not found on %s p%d: %r" % (os.path.basename(pdf), idx, s_))
        r = hits[0]
        out.append((r.x0 / w, (r.y0 + r.y1) / 2 / h))
    ratio = h / w
    d.close()
    return out, ratio


def detail(p, png, crop, pins, feats, ratio=11 / 8.5):
    """One page up close. crop = (x0, y0, share of the page width that fills the frame);
    pins = [(x, y)] as fractions of the page, from find_pins."""
    x0, y0, zoom_w = crop
    img_w = 640 / zoom_w
    img_h = img_w * ratio
    left, top = -x0 * img_w, -y0 * img_h
    dots = ""
    for n, (px, py) in enumerate(pins, 1):
        fx = 70 + (px - x0) * img_w - 34
        fy = 70 + (py - y0) * img_h
        dots += '<div class="pin" style="left:%dpx;top:%dpx">%d</div>' % (max(58, fx) - 20, fy - 20, n)
    fs = "".join('<div><b>%d</b><section><h3>%s</h3><p>%s</p></section></div>' % (i + 1, esc(a), esc(b))
                 for i, (a, b) in enumerate(feats))
    return art('<div class="crop"><img src="%s" style="width:%dpx;left:%dpx;top:%dpx"></div>%s'
               '<div class="feat">%s</div>' % (b64(png), img_w, left, top, dots, fs))


def features_only(p, screen_html, feats):
    fs = "".join('<div><b>%d</b><section><h3>%s</h3><p>%s</p></section></div>' % (i + 1, esc(a), esc(b))
                 for i, (a, b) in enumerate(feats))
    return art('<div class="crop" style="width:660px">%s</div><div class="feat">%s</div>' % (screen_html, fs))


def use_ipad(p, png, second):
    inner = (brandline("How you use it") + '<div class="copy" style="top:170px;width:440px"><h1 style="font-size:64px">%s</h1>'
             '<p class="sub">%s</p></div>' % (rich(p["use_title"]), esc(p["use_sub"])) +
             '<div class="sheet" style="left:880px;top:150px;width:330px;height:427px;transform:rotate(6deg)"><img src="%s"></div>' % b64(second) +
             '<div class="ipad" style="left:560px;top:250px;width:520px;height:690px;transform:rotate(-3deg)"><div class="scr"><img src="%s"></div></div>' % b64(png))
    return art(inner)


def use_phones(p, phone_pngs):
    ph = ""
    spots = [(600, 230, -6), (810, 170, 0), (1010, 240, 6)]
    for (x, y, rot), png in zip(spots, phone_pngs):
        ph += ('<div class="phone" style="left:%dpx;top:%dpx;width:250px;height:520px;transform:rotate(%sdeg)">'
               '<div class="scr"><img src="%s"></div></div>' % (x - 120, y, rot, b64(png)))
    inner = (brandline("How you use it") + '<div class="copy" style="top:170px;width:390px"><h1 style="font-size:60px">%s</h1>'
             '<p class="sub">%s</p></div>' % (rich(p["use_title"]), esc(p["use_sub"])) + ph)
    return art(inner)


def use_texts(p, lines):
    bubbles = "".join('<div style="align-self:%s;max-width:82%%;background:%s;color:%s;border-radius:20px;padding:10px 14px;'
                      'font-size:13.5px;line-height:1.35">%s</div>' % (
                          "flex-end" if i % 2 else "flex-start", "#a86b7e" if i % 2 else "#efe9ea",
                          "#fff" if i % 2 else "#2a2326", esc(t)) for i, t in enumerate(lines))
    phone = ('<div class="phone" style="left:700px;top:120px;width:360px;height:760px;transform:rotate(3deg)"><div class="scr" '
             'style="display:flex;flex-direction:column;gap:10px;padding:70px 16px 16px;background:#fff;font-family:var(--sans)">'
             '%s</div></div>' % bubbles)
    inner = (brandline("How you use it") + '<div class="copy" style="top:170px;width:440px"><h1 style="font-size:64px">%s</h1>'
             '<p class="sub">%s</p></div>' % (rich(p["use_title"]), esc(p["use_sub"])) + phone)
    return art(inner)


def use_print(p, pngs):
    inner = (brandline("How you use it") + '<div class="copy" style="top:170px;width:440px"><h1 style="font-size:64px">%s</h1>'
             '<p class="sub">%s</p></div>' % (rich(p["use_title"]), esc(p["use_sub"])) +
             '<div class="sheet" style="left:600px;top:130px;width:420px;height:543px;transform:rotate(-5deg)"><img src="%s"></div>' % b64(pngs[0]) +
             '<div class="sheet" style="left:790px;top:330px;width:400px;height:518px;transform:rotate(4deg)"><img src="%s"></div>' % b64(pngs[1]))
    return art(inner)


def use_laptop(p, screen):
    inner = (brandline("How you use it") + '<div class="top" style="top:120px"><h1 style="font-size:58px">Opens in Excel, <i>Google Sheets</i> and Numbers</h1></div>'
             '<div class="laptop" style="left:195px;top:290px"><div class="lid"><div class="scr" style="height:540px">%s</div></div><div class="base"></div></div>' % screen)
    return art(inner)


def maker(p, formats):
    promises = ["Designed by Evelyn Romero, licensed esthetician, from the tools she uses at Lumevina Aesthetics in "
                "Woodland Hills, California.", "Instant download. Nothing to wait for, nothing shipped.", formats,
                "Questions? Message us. We reply within one business day."]
    return ('<div class="art"><div class="maker"><div class="pic"><img src="%s"></div><div class="txt">%s'
            '<h1>Made in a <i>real studio</i></h1><div class="promise">%s</div></div></div></div>'
            % (b64(os.path.join(ART, "mark-rose.png")), brandline("Lumevina Aesthetics"),
               "".join("<div>%s</div>" % esc(x) for x in promises)))


FORMATS = {"01": "Fillable PDF and Word, in US Letter and A4.", "02": "Print sheets, phone images and Word.",
           "03": "Fillable PDF and Word, in US Letter and A4.", "04": "Word, Google Docs or Pages, plus sample PDFs.",
           "05": "An Excel file that also opens in Google Sheets and Numbers.",
           "06": "An Excel file that also opens in Google Sheets and Numbers.",
           "07": "PDF and Word, in US Letter and A4.", "08": "Fillable PDF in US Letter and A4.",
           "09": "Fillable PDF in US Letter and A4.", "10": "PDF, Word and Excel files in one download."}


def build(mods, dist, tmp, render, only=None):
    import listings
    P = plans(dist, tmp)
    jobs, outs = [], []
    for m in mods:
        key = m.SLUG[:2]
        if only and m.SLUG not in only:
            continue
        p = P[key]
        work = os.path.join(tmp, m.SLUG, "mock")
        os.makedirs(work, exist_ok=True)
        imgdir = os.path.join(dist, m.SLUG, "images")
        shutil.rmtree(imgdir, ignore_errors=True)
        os.makedirs(imgdir)
        price = listings.CATALOG[key]["price"]
        pg = lambda pdf, i, dpi=150: page_png(pdf, i, os.path.join(work, "%s-%d-%d.png" % (
            os.path.basename(pdf)[:12].replace(" ", "_"), i, dpi)), dpi)
        arts = []
        if p.get("bundle"):
            covers = []
            for mm in mods[:9]:
                pp = P[mm.SLUG[:2]]
                if "pdf" in pp:
                    covers.append((pg(pp["pdf"], pp["fan"][1] if len(pp["fan"]) > 1 else 0), mm.NAME,
                                   "11/8.5" if pp.get("landscape") else "8.5/11"))
                else:
                    covers.append((None, mm.NAME, "8.5/11", pp))
            first = [c[0] for c in covers if c[0]][:3]
            arts.append(cover(p, [first[1], first[0], first[2]], price))
            tiles = []
            for c in covers:
                if c[0]:
                    tiles.append((c[0], c[1], c[2]))
            arts.append(inside(p, tiles, title="Nine products, <i>one kit</i>"))
            xl = sheet_html(P["05"]["xlsx"], "Calculator", 30, 4, P["05"]["tabs"])
            arts.append(features_only(p, xl, [
                ("Forms and cards", "Intake, consent, analysis and aftercare: everything clients see and sign."),
                ("Menu and scripts", "What you offer, and the words that get clients to rebook and join."),
                ("Numbers and plan", "Membership pricing, bookkeeping and a 90-day plan to open strong.")]))
            arts.append(use_ipad(dict(p, use_title="Everything for _day one_",
                                      use_sub="Open the kit, set up your forms and menu, and see your first client "
                                              "the same week."), first[0], first[1]))
        elif "xlsx" in p:
            screens = [sheet_html(p["xlsx"], n, r, c, p["tabs"]) for n, r, c in p["sheets"]]
            hero = ('<div class="laptop" style="left:585px;top:250px;width:640px"><div class="lid"><div class="scr" style="height:430px">%s</div></div><div class="base"></div></div>' % screens[0])
            arts.append(art(brandline() + '<div class="copy"><h1>%s</h1><p class="sub">%s</p>%s</div>' % (
                rich(p["headline"]), esc(p["sub"]), chips(p["chips"])) + hero +
                '<div class="price"><b>Instant download</b>edit, print, use today</div>'))
            tabs_tiles = "".join('<div class="tile"><div class="sheet" style="height:330px;overflow:hidden">%s</div><p>%s</p></div>'
                                 % (s, esc(n)) for s, (n, _, _) in zip(screens, p["sheets"]))
            arts.append(art(brandline("What's inside") + '<div class="top"><h1>What&rsquo;s <i>inside</i></h1></div>'
                            '<div class="grid" style="grid-template-columns:repeat(3,1fr);top:300px">%s</div>' % tabs_tiles))
            arts.append(features_only(p, screens[0], p["detail_feats"]))
            arts.append(use_laptop(p, screens[1]))
        else:
            pdf = p["pdf"]
            pngs = [pg(pdf, i) for i in p["fan"]]
            arts.append(cover(p, pngs, price))
            if p["inside"] == "phone":
                phones = [os.path.join(tmp, m.SLUG, "phone-%d.png" % (i + 1)) for i in range(8)]
                import p02_aftercare
                arts.append(inside(p, [(ph, c[0], "9/16") for ph, c in zip(phones, p02_aftercare.CARDS)]))
            else:
                tiles = [(pg(a, b), c, "8.5/11") for a, b, c in p["inside"]]
                arts.append(inside(p, tiles))
            if p.get("detail"):
                dpdf, di, crop, texts, feats = p["detail"]
                pins, ratio = find_pins(dpdf, di, texts)
                arts.append(detail(p, pg(dpdf, di, 220), crop, pins, feats, ratio))
            else:
                # the aftercare set: one card of a print sheet up close
                pins, ratio = find_pins(pdf, 1, ["The first 48 hours", "What's normal", "Call me if"])
                arts.append(detail(p, pg(pdf, 1, 220), (0.02, 0.06, 0.52), pins,
                                   [("The first day", "Exactly what to skip and what to do, in plain words."),
                                    ("What's normal", "So clients don't panic at a little pinkness or flaking."),
                                    ("When to call", "Clear signs to reach out, and room for your number.")], ratio))
            if p["use"] == "ipad":
                arts.append(use_ipad(p, pngs[1] if len(pngs) > 1 else pngs[0], pngs[0]))
            elif p["use"] == "phones":
                arts.append(use_phones(p, [os.path.join(tmp, m.SLUG, "phone-%d.png" % i) for i in (1, 3, 7)]))
            elif p["use"] == "phones_text":
                arts.append(use_texts(p, ["You're booked! Signature Facial, Tue, Mar 14 at 10:00. Reply C to confirm.",
                                          "C", "Hi Maria! How's your skin feeling after your facial? A little "
                                          "pinkness is normal. Anything else, just text me.",
                                          "It's glowing!! Thank you", "Four weeks already! I have Tue at 10 or Sat "
                                          "at 1. Want one?", "Sat at 1 please"]))
            elif p["use"] == "print":
                arts.append(use_print(p, [pg(pdf, 0), pg(pdf, 2)]))
        arts.append(maker(p, FORMATS[key]))
        src = os.path.join(work, "mock.html")
        open(src, "w").write(doc(arts))
        jobs.append({"kind": "shot", "html": src, "out": os.path.join(imgdir, "%s" % key), "selector": ".art",
                     "width": W, "height": H, "scale": 2, "type": "jpeg", "quality": 86, "wait": 400})
    # the shop's own art
    shop = os.path.join(dist, "etsy", "shop")
    os.makedirs(shop, exist_ok=True)
    src = os.path.join(tmp, "shop.html")
    open(src, "w").write(shop_html())
    from PIL import Image
    Image.open(os.path.join(HERE, "..", "..", "img", "app-icon-1024.png")).convert("RGB").resize(
        (500, 500), Image.LANCZOS).save(os.path.join(shop, "icon.png"))
    jobs.append({"kind": "shot", "html": src, "out": os.path.join(shop, "banner.jpg"), "selector": ".banner", "width": 1680,
                 "height": 420, "scale": 2, "type": "jpeg", "quality": 90, "single": True})
    rep = render(jobs)
    for r in rep:
        if r.get("errors"):
            print("image errors", r)
    return rep


def shop_html():
    mark = b64(os.path.join(ART, "mark-rose.png"))
    css = """
.banner{width:1680px;height:420px;position:relative;overflow:hidden;
  background:radial-gradient(120% 140% at 80% 30%,#fbf7f7 0%,#f1e6e9 60%,#e9dce0 100%);display:flex;align-items:center}
.banner img{position:absolute;right:120px;top:40px;height:360px;opacity:.85}
.banner .t{margin-left:120px}
.banner h1{font-family:var(--serif);font-weight:500;font-size:70px;letter-spacing:.2em;margin:0;line-height:1}
.banner p{margin:18px 0 0;font-size:19px;letter-spacing:.28em;text-transform:uppercase;color:var(--rose)}
.banner .s{margin-top:22px;font-size:20px;color:var(--muted);font-weight:300;letter-spacing:0}
"""
    return ('<!doctype html><html><head><meta charset="utf-8"><style>%s%s%s</style></head><body>'
            '<div class="banner"><div class="t"><h1>LUMEVINA</h1><p>Studio · templates and tools</p>'
            '<div class="s">Made by a licensed esthetician, from a real studio in Los Angeles.</div></div>'
            '<img src="%s"></div></body></html>' % (font_css(), TOKENS, css, mark))
