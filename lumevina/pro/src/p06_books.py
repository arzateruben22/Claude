"""06 · Bookkeeping & Profit Tracker: income, expenses, products and a month-by-month profit picture."""
import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment

import sheets as S
from p05_membership import start_sheet

SLUG = "06-bookkeeping-profit-tracker"
NAME = "Bookkeeping & Profit Tracker"

ROWS = 600                      # rows ready in Income and Expenses
FIRST = 6                       # first data row in Income and Expenses
LAST = FIRST + ROWS - 1

INCOME_CATS = ["Facials", "Peels and advanced", "Waxing", "Lash and brow", "Memberships", "Products",
               "Gift cards sold", "Other services"]
EXPENSE_CATS = ["Rent and utilities", "Backbar supplies", "Retail inventory", "Equipment", "Insurance and licenses",
                "Software and booking", "Marketing", "Education", "Laundry and cleaning", "Bank fees",
                "Phone and internet", "Car and mileage", "Other"]
PAY = ["Card", "Cash", "Venmo or Zelle", "Gift card", "Other"]
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def settings(ws):
    S.title(ws, "Settings", "Set these once. The lists feed the drop-downs on the other tabs.")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 4
    ws.column_dimensions["D"].width = 28
    ws.column_dimensions["E"].width = 4
    ws.column_dimensions["F"].width = 22
    S.head(ws, 4, "Your business", 6)
    S.label(ws, "A5", "Business name")
    S.inp(ws, "B5", "Your Studio Name")
    S.label(ws, "A6", "Year")
    S.inp(ws, "B6", 2026, "0", key=True)
    S.label(ws, "A7", "Set aside for taxes")
    S.inp(ws, "B7", 0.25, S.PCT, key=True)
    S.note(ws, "D7", "Ask your tax preparer for your number.")
    S.label(ws, "A8", "Card fee, percent")
    S.inp(ws, "B8", 0.029, S.PCT1)
    S.label(ws, "A9", "Card fee, per charge")
    S.inp(ws, "B9", 0.30, S.USD2)
    S.note(ws, "D8", "Your processor's rate (Square, Stripe, etc.).")
    for col, title, items in (("A", "Income categories", INCOME_CATS), ("D", "Expense categories", EXPENSE_CATS),
                              ("F", "Ways to pay", PAY)):
        c = ws["%s11" % col]
        c.value = title.upper()
        c.font = S.HEAD
        for i, it in enumerate(items):
            S.inp(ws, "%s%d" % (col, 12 + i), it)
    S.note(ws, "A26", "Rename categories to fit your studio; keep the same number of rows so the totals stay linked.")


def income(ws):
    S.title(ws, "Income", "One line per sale. Tips and card fees are worked out for you.")
    cols = [("Date", 12), ("Client", 20), ("What", 26), ("Category", 20), ("Amount", 11), ("Tip", 9),
            ("Paid by", 15), ("Card fee", 10), ("You keep", 11), ("Month", 8)]
    for i, (h, w) in enumerate(cols):
        c = ws.cell(row=5, column=i + 1, value=h.upper())
        c.font = S.HEAD
        c.border = S.LINE
        ws.column_dimensions["ABCDEFGHIJ"[i]].width = w
    S.note(ws, "A3", "Example on the first line: replace it with your first sale. Blue columns are yours; black ones fill themselves in.")
    example = [datetime.date(2026, 1, 6), "Maria G.", "Signature Facial", "Facials", 125, 20, "Card"]
    for r in range(FIRST, LAST + 1):
        vals = example if r == FIRST else [None] * 7
        for i, v in enumerate(vals):
            c = ws.cell(row=r, column=i + 1, value=v)
            c.font = S.INPUT
        ws.cell(row=r, column=1).number_format = "mm/dd/yyyy"
        ws.cell(row=r, column=5).number_format = S.USD2
        ws.cell(row=r, column=6).number_format = S.USD2
        S.calc(ws, "H%d" % r, '=IF(G{r}="Card",ROUND((E{r}+F{r})*Settings!$B$8+Settings!$B$9,2),0)'.format(r=r), S.USD2)
        S.calc(ws, "I%d" % r, '=IF(A{r}="","",E{r}+F{r}-H{r})'.format(r=r), S.USD2)
        S.calc(ws, "J%d" % r, '=IF(A{r}="","",MONTH(A{r}))'.format(r=r), "0")
    S.dropdown(ws, "D%d:D%d" % (FIRST, LAST), "=Settings!$A$12:$A$19")
    S.dropdown(ws, "G%d:G%d" % (FIRST, LAST), "=Settings!$F$12:$F$16")
    ws.freeze_panes = "A6"


def expenses(ws):
    S.title(ws, "Expenses", "Everything the business pays for. Keep the receipts: your tax preparer will ask.")
    cols = [("Date", 12), ("Paid to", 22), ("What", 30), ("Category", 22), ("Amount", 11), ("Paid by", 15),
            ("Receipt saved", 13), ("Month", 8)]
    for i, (h, w) in enumerate(cols):
        c = ws.cell(row=5, column=i + 1, value=h.upper())
        c.font = S.HEAD
        c.border = S.LINE
        ws.column_dimensions["ABCDEFGH"[i]].width = w
    S.note(ws, "A3", "Example on the first line: replace it with your first expense.")
    example = [datetime.date(2026, 1, 2), "Skin supply co.", "Cleanser and masks (backbar)", "Backbar supplies",
               86.4, "Card", "Yes"]
    for r in range(FIRST, LAST + 1):
        vals = example if r == FIRST else [None] * 7
        for i, v in enumerate(vals):
            ws.cell(row=r, column=i + 1, value=v).font = S.INPUT
        ws.cell(row=r, column=1).number_format = "mm/dd/yyyy"
        ws.cell(row=r, column=5).number_format = S.USD2
        S.calc(ws, "H%d" % r, '=IF(A{r}="","",MONTH(A{r}))'.format(r=r), "0")
    S.dropdown(ws, "D%d:D%d" % (FIRST, LAST), "=Settings!$D$12:$D$24")
    S.dropdown(ws, "F%d:F%d" % (FIRST, LAST), "=Settings!$F$12:$F$16")
    S.dropdown(ws, "G%d:G%d" % (FIRST, LAST), ["Yes", "No"])
    ws.freeze_panes = "A6"


def products(ws):
    S.title(ws, "Products", "Your shelf: what each product earns, what's in stock and what to reorder.")
    cols = [("Product", 28), ("Brand", 16), ("Your cost", 11), ("Price", 11), ("Profit each", 12), ("Margin", 9),
            ("In stock", 10), ("Reorder at", 11), ("Reorder?", 11), ("Stock value", 12)]
    for i, (h, w) in enumerate(cols):
        c = ws.cell(row=5, column=i + 1, value=h.upper())
        c.font = S.HEAD
        c.border = S.LINE
        ws.column_dimensions["ABCDEFGHIJ"[i]].width = w
    example = ["Gentle gel cleanser", "Your brand", 14, 32, None, None, 6, 3]
    for r in range(6, 106):
        vals = example if r == 6 else [None] * 8
        for i, v in enumerate(vals[:4] + [None, None] + vals[6:8]):
            if i in (4, 5):
                continue
            ws.cell(row=r, column=i + 1, value=v).font = S.INPUT
        for col in (3, 4):
            ws.cell(row=r, column=col).number_format = S.USD2
        S.calc(ws, "E%d" % r, '=IF(D{r}="","",D{r}-C{r})'.format(r=r), S.USD2)
        S.calc(ws, "F%d" % r, '=IF(OR(D{r}="",D{r}=0),"",(D{r}-C{r})/D{r})'.format(r=r), S.PCT)
        S.calc(ws, "I%d" % r, '=IF(OR(G{r}="",H{r}=""),"",IF(G{r}<=H{r},"Reorder",""))'.format(r=r))
        S.calc(ws, "J%d" % r, '=IF(OR(C{r}="",G{r}=""),"",C{r}*G{r})'.format(r=r), S.USD2)
    S.label(ws, "A108", "Total stock value", bold=True)
    S.calc(ws, "J108", "=SUM(J6:J105)", S.USD2, bold=True)
    S.note(ws, "A109", "Aim for a margin of 50% or more on retail. Under 40%, check your price or your supplier.")
    ws.freeze_panes = "A6"


def monthly(ws):
    S.title(ws, "Month by Month", "Filled in from your Income and Expenses tabs. Nothing to type here.")
    ws.column_dimensions["A"].width = 26
    for i in range(13):
        ws.column_dimensions["BCDEFGHIJKLMN"[i]].width = 11
    for i, m in enumerate(MONTHS + ["Year"]):
        c = ws.cell(row=4, column=2 + i, value=m.upper())
        c.font = S.HEAD
        c.alignment = Alignment(horizontal="right")
        c.border = S.LINE
    ws.cell(row=4, column=1).border = S.LINE
    inc = "Income!$E$%d:$E$%d" % (FIRST, LAST)
    icat = "Income!$D$%d:$D$%d" % (FIRST, LAST)
    imon = "Income!$J$%d:$J$%d" % (FIRST, LAST)
    iyear = "Income!$A$%d:$A$%d" % (FIRST, LAST)
    r = 5
    S.label(ws, "A%d" % r, "INCOME", bold=True)
    r += 1
    first_inc = r
    for i in range(len(INCOME_CATS)):
        S.calc(ws, "A%d" % r, "=Settings!A%d" % (12 + i)).font = S.LINK
        for m in range(12):
            col = "BCDEFGHIJKLM"[m]
            S.calc(ws, "%s%d" % (col, r), '=SUMIFS(%s,%s,$A%d,%s,%d,%s,">="&DATE(Settings!$B$6,1,1),%s,"<="&DATE(Settings!$B$6,12,31))'
                   % (inc, icat, r, imon, m + 1, iyear, iyear), S.USD)
        S.calc(ws, "N%d" % r, "=SUM(B%d:M%d)" % (r, r), S.USD, bold=True)
        r += 1
    last_inc = r - 1
    S.label(ws, "A%d" % r, "Tips")
    for m in range(12):
        col = "BCDEFGHIJKLM"[m]
        S.calc(ws, "%s%d" % (col, r), '=SUMIFS(Income!$F$%d:$F$%d,%s,%d,%s,">="&DATE(Settings!$B$6,1,1),%s,"<="&DATE(Settings!$B$6,12,31))'
               % (FIRST, LAST, imon, m + 1, iyear, iyear), S.USD)
    S.calc(ws, "N%d" % r, "=SUM(B%d:M%d)" % (r, r), S.USD, bold=True)
    tips = r
    r += 1
    S.label(ws, "A%d" % r, "Total income", bold=True)
    for col in "BCDEFGHIJKLMN":
        S.calc(ws, "%s%d" % (col, r), "=SUM(%s%d:%s%d)" % (col, first_inc, col, tips), S.USD, bold=True)
        ws["%s%d" % (col, r)].border = S.TOPLINE
    tot_inc = r
    r += 2
    S.label(ws, "A%d" % r, "EXPENSES", bold=True)
    r += 1
    S.label(ws, "A%d" % r, "Card fees")
    for m in range(12):
        col = "BCDEFGHIJKLM"[m]
        S.calc(ws, "%s%d" % (col, r), '=SUMIFS(Income!$H$%d:$H$%d,%s,%d,%s,">="&DATE(Settings!$B$6,1,1),%s,"<="&DATE(Settings!$B$6,12,31))'
               % (FIRST, LAST, imon, m + 1, iyear, iyear), S.USD)
    S.calc(ws, "N%d" % r, "=SUM(B%d:M%d)" % (r, r), S.USD, bold=True)
    first_exp = r
    r += 1
    exa = "Expenses!$E$%d:$E$%d" % (FIRST, LAST)
    ecat = "Expenses!$D$%d:$D$%d" % (FIRST, LAST)
    emon = "Expenses!$H$%d:$H$%d" % (FIRST, LAST)
    eyear = "Expenses!$A$%d:$A$%d" % (FIRST, LAST)
    for i in range(len(EXPENSE_CATS)):
        S.calc(ws, "A%d" % r, "=Settings!D%d" % (12 + i)).font = S.LINK
        for m in range(12):
            col = "BCDEFGHIJKLM"[m]
            S.calc(ws, "%s%d" % (col, r), '=SUMIFS(%s,%s,$A%d,%s,%d,%s,">="&DATE(Settings!$B$6,1,1),%s,"<="&DATE(Settings!$B$6,12,31))'
                   % (exa, ecat, r, emon, m + 1, eyear, eyear), S.USD)
        S.calc(ws, "N%d" % r, "=SUM(B%d:M%d)" % (r, r), S.USD, bold=True)
        r += 1
    last_exp = r - 1
    S.label(ws, "A%d" % r, "Total expenses", bold=True)
    for col in "BCDEFGHIJKLMN":
        S.calc(ws, "%s%d" % (col, r), "=SUM(%s%d:%s%d)" % (col, first_exp, col, last_exp), S.USD, bold=True)
        ws["%s%d" % (col, r)].border = S.TOPLINE
    tot_exp = r
    r += 2
    S.label(ws, "A%d" % r, "Profit", bold=True)
    for col in "BCDEFGHIJKLMN":
        S.calc(ws, "%s%d" % (col, r), "=%s%d-%s%d" % (col, tot_inc, col, tot_exp), S.USD, bold=True)
        ws["%s%d" % (col, r)].fill = S.BAND
    ws["A%d" % r].fill = S.BAND
    profit = r
    r += 1
    S.label(ws, "A%d" % r, "Set aside for taxes")
    for col in "BCDEFGHIJKLMN":
        S.calc(ws, "%s%d" % (col, r), "=MAX(0,%s%d)*Settings!$B$7" % (col, profit), S.USD)
    tax = r
    r += 1
    S.label(ws, "A%d" % r, "Yours to keep", bold=True)
    for col in "BCDEFGHIJKLMN":
        S.calc(ws, "%s%d" % (col, r), "=%s%d-%s%d" % (col, profit, col, tax), S.USD, bold=True)
    keep = r
    r += 1
    S.label(ws, "A%d" % r, "Profit margin")
    for col in "BCDEFGHIJKLMN":
        S.calc(ws, "%s%d" % (col, r), "=IF(%s%d=0,0,%s%d/%s%d)" % (col, tot_inc, col, profit, col, tot_inc), S.PCT)
    ws.freeze_panes = "B5"
    return {"tot_inc": tot_inc, "tot_exp": tot_exp, "profit": profit, "keep": keep, "tax": tax,
            "first_exp": first_exp, "last_exp": last_exp, "first_inc": first_inc, "last_inc": last_inc}


def glance(ws, rows):
    S.title(ws, "The Year at a Glance", "The numbers that matter, all in one place.")
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 40
    S.head(ws, 4, "This year", 3)
    items = [
        ("Total income", "='Month by Month'!N%d" % rows["tot_inc"], S.USD),
        ("Total expenses", "='Month by Month'!N%d" % rows["tot_exp"], S.USD),
        ("Profit", "='Month by Month'!N%d" % rows["profit"], S.USD),
        ("Set aside for taxes", "='Month by Month'!N%d" % rows["tax"], S.USD),
        ("Yours to keep", "='Month by Month'!N%d" % rows["keep"], S.USD),
        ("Profit margin", "=IF(B5=0,0,B7/B5)", S.PCT),
        ("Average profit a month so far", "=IF(COUNTIF('Month by Month'!B%d:M%d,\"<>0\")=0,0,B7/COUNTIF('Month by Month'!B%d:M%d,\"<>0\"))"
         % (rows["tot_inc"], rows["tot_inc"], rows["tot_inc"], rows["tot_inc"]), S.USD),
        ("Best month", "=IF(MAX('Month by Month'!B%d:M%d)=0,\"-\",INDEX('Month by Month'!B4:M4,MATCH(MAX('Month by Month'!B%d:M%d),'Month by Month'!B%d:M%d,0)))"
         % (rows["profit"], rows["profit"], rows["profit"], rows["profit"], rows["profit"], rows["profit"]), None),
        ("Biggest income category", "=IF(MAX('Month by Month'!N%d:N%d)=0,\"-\",INDEX('Month by Month'!A%d:A%d,MATCH(MAX('Month by Month'!N%d:N%d),'Month by Month'!N%d:N%d,0)))"
         % tuple([rows["first_inc"], rows["last_inc"]] * 4), None),
        ("Biggest expense category", "=IF(MAX('Month by Month'!N%d:N%d)=0,\"-\",INDEX('Month by Month'!A%d:A%d,MATCH(MAX('Month by Month'!N%d:N%d),'Month by Month'!N%d:N%d,0)))"
         % tuple([rows["first_exp"], rows["last_exp"]] * 4), None),
        ("Retail stock on the shelf", "=Products!J108", S.USD),
    ]
    for i, (lab, f, fmt) in enumerate(items):
        r = 5 + i
        S.label(ws, "A%d" % r, lab, bold=lab in ("Profit", "Yours to keep"))
        c = S.calc(ws, "B%d" % r, f, fmt, bold=lab in ("Profit", "Yours to keep"))
        c.alignment = Alignment(horizontal="right")
        ws["A%d" % r].border = S.LINE
        ws["B%d" % r].border = S.LINE
    S.note(ws, "C8", "A cushion for income tax and self-employment tax.")
    S.note(ws, "C18", "Not a tax return. Share this workbook with your tax preparer at year end.")


def make(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Start Here"
    start_sheet(ws, NAME, [
        "Open Settings. Put in the year, your card fee and how much to set aside for taxes. Rename the categories if you like.",
        "Every sale goes on the Income tab: date, client, what, category, amount, tip and how they paid.",
        "Every business cost goes on the Expenses tab, with its category. Save the receipt.",
        "List your retail products on the Products tab to see margins and when to reorder.",
        "Month by Month and The Year at a Glance fill themselves in. Check them on the first of each month.",
    ], [("Good to know", "The first line of Income and Expenses is an example. Type over it. This tracker helps you "
                         "stay organized and is not tax advice; share it with your tax preparer."),
        ("Made by", "Evelyn Romero, licensed esthetician, Lumevina Aesthetics. Questions? Message us on Etsy.")])
    settings(wb.create_sheet("Settings"))
    income(wb.create_sheet("Income"))
    expenses(wb.create_sheet("Expenses"))
    products(wb.create_sheet("Products"))
    rows = monthly(wb.create_sheet("Month by Month"))
    glance(wb.create_sheet("The Year at a Glance"), rows)
    for w in wb.worksheets:
        w.sheet_properties.tabColor = S.ROSE
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return S.recalc(path)


START = {
    "inside": [
        ("Spreadsheet", "The tracker as an Excel file. Opens in Excel, Google Sheets (File, Import) and Numbers."),
        ("Seven tabs", "Start Here · Settings · Income · Expenses · Products · Month by Month · The Year at a Glance"),
    ],
    "tips": [
        "Set the year, your card fee and your tax set-aside on the Settings tab.",
        "Log each sale on Income and each cost on Expenses. Five minutes at the end of the day is plenty.",
        "Add your retail products once, and update the stock when you count the shelf.",
        "On the first of each month, look at Month by Month: what grew, what cost more, and what you kept.",
    ],
    "extra": [("h", "Google Sheets"),
              ("p", "In Google Drive, choose **New, File upload**, pick the file, then **Open with Google Sheets**. "
                    "The drop-downs and totals work the same way."),
              ("h", "At tax time"),
              ("p", "Share the workbook, or the Month by Month and Expenses tabs, with your tax preparer. The "
                    "categories line up with the way most preparers group business costs.")],
    "legal": False,
    "fillable": False,
}


def build(c):
    fn = "Bookkeeping & Profit Tracker.xlsx"
    c.after.append(lambda: make(c.path(fn)))
    c.add(fn)
    c.start_here(formats=())
