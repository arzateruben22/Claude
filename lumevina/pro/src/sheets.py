"""Shared look for the spreadsheets: Arial, blue inputs, black formulas, yellow key inputs."""
import json, os, subprocess, sys

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROSE = "A86B7E"
INKC = "2A2326"
MUTEDC = "7A6D72"
BLUE = "0000FF"
BLUSH = "F7EEF1"
YELLOW = "FFF2A8"
HAIRC = "E7DDE0"

F = "Arial"
TITLE = Font(name=F, size=18, bold=True, color=INKC)
SUB = Font(name=F, size=10, color=MUTEDC, italic=True)
HEAD = Font(name=F, size=9, bold=True, color=ROSE)
LABEL = Font(name=F, size=10, color=INKC)
BOLD = Font(name=F, size=10, bold=True, color=INKC)
INPUT = Font(name=F, size=10, color=BLUE)
CALC = Font(name=F, size=10, color="000000")
BIG = Font(name=F, size=14, bold=True, color="000000")
NOTE = Font(name=F, size=8.5, color=MUTEDC, italic=True)
LINK = Font(name=F, size=10, color="008000")

KEY = PatternFill("solid", fgColor=YELLOW)
BAND = PatternFill("solid", fgColor=BLUSH)
LINE = Border(bottom=Side(style="thin", color=HAIRC))
TOPLINE = Border(top=Side(style="thin", color=ROSE))

USD = '$#,##0;($#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
PCT = '0%;(0%);"-"'
PCT1 = '0.0%;(0.0%);"-"'
NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'

SKILL_RECALC = os.environ.get("RECALC", "/root/.claude/skills/synced/ecf5ee15-05b7-4201-9b10-f52a4f28a4c9_410bbf28-206a-419c-bdef-c2df35665e27/xlsx/scripts/recalc.py")


def title(ws, text, sub, width=None):
    ws["A1"] = text
    ws["A1"].font = TITLE
    ws["A2"] = sub
    ws["A2"].font = SUB
    ws.row_dimensions[1].height = 26
    ws.sheet_view.showGridLines = False


def head(ws, row, text, cols=6):
    c = ws.cell(row=row, column=1, value=text.upper())
    c.font = HEAD
    for col in range(1, cols + 1):
        ws.cell(row=row, column=col).border = LINE
    ws.row_dimensions[row].height = 22


def inp(ws, ref, value, fmt=None, key=False):
    c = ws[ref]
    c.value = value
    c.font = INPUT
    if fmt:
        c.number_format = fmt
    if key:
        c.fill = KEY
    return c


def calc(ws, ref, formula, fmt=None, bold=False, big=False):
    c = ws[ref]
    c.value = formula
    c.font = BIG if big else (Font(name=F, size=10, bold=True, color="000000") if bold else CALC)
    if fmt:
        c.number_format = fmt
    return c


def label(ws, ref, text, bold=False):
    c = ws[ref]
    c.value = text
    c.font = BOLD if bold else LABEL
    c.alignment = Alignment(vertical="center", wrap_text=False)
    return c


def note(ws, ref, text):
    c = ws[ref]
    c.value = text
    c.font = NOTE
    c.alignment = Alignment(vertical="center", wrap_text=False)
    return c


def legend(ws, row, col=1):
    """The three-line key that says which cells to change."""
    items = [("Blue text", "Your numbers. Change these.", INPUT, None),
             ("Yellow cell", "The most important inputs.", INPUT, KEY),
             ("Black text", "Calculated for you. Leave these alone.", CALC, None)]
    for i, (a, b, font, fill) in enumerate(items):
        c = ws.cell(row=row + i, column=col, value=a)
        c.font = font
        if fill:
            c.fill = fill
        ws.cell(row=row + i, column=col + 1, value=b).font = LABEL


def dropdown(ws, rng, options_ref_or_list):
    if isinstance(options_ref_or_list, list):
        formula = '"%s"' % ",".join(options_ref_or_list)
    else:
        formula = options_ref_or_list
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng)


def recalc(path):
    """Calculate every formula with LibreOffice so the file opens with values everywhere."""
    if not os.path.exists(SKILL_RECALC):
        print("  (recalc skipped: set RECALC to recalc.py; Excel and Google Sheets calculate on open)")
        return None
    r = subprocess.run([sys.executable, SKILL_RECALC, path, "90"], capture_output=True, text=True)
    try:
        out = json.loads(r.stdout)
    except ValueError:
        sys.exit("recalc failed: " + r.stdout + r.stderr)
    if out.get("status") != "success":
        sys.exit("formula problems in %s: %s" % (path, json.dumps(out)[:800]))
    return out
