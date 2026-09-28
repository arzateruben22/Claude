"""One content source, two outputs: print-ready HTML (then PDF) and an editable Word file.

A document is a dict:
    {"title": ..., "header": "studio" | "brand" | None, "foot": "..." ,
     "pages": [{"name": ..., "blocks": [...]}, ...]}

Blocks are tuples; the first item names the kind:
    ("title", text, subtitle)            form or chapter title
    ("h", text)                          section heading with a hairline
    ("p", text)                          paragraph; **bold** and _italic_ work
    ("small", text)                      fine print
    ("list", [items])                    bullets
    ("steps", [items])                   numbered list (a real sequence)
    ("fields", [(label, weight, name)])  a row of write-on lines
    ("lines", n, name[, label])          n ruled lines (one multi-line field)
    ("checks", label, [options], name, cols)
    ("yn", [(question, name)])           yes / no / details rows
    ("initial", text, name)              a statement the client initials
    ("sign", [(label, weight, name)])    signature row
    ("table", headers, widths, rows, name[, row_h])  rows: int (blank) or lists
    ("svg", markup, height_in, caption)
    ("cols", left_blocks, right_blocks, left_share)
    ("script", when, words)              words to say, set as a quote
    ("callout", title, text)
    ("box", label, height_in, name)
    ("kv", [(term, text)])
    ("space", inches)

Anything with a name becomes a fillable field in the PDF (see render.cjs and
fill.py): text lines, checkboxes, table cells and boxes.
"""
import html as _h, re

from brand import (TOKENS, font_css, INK, MUTED, RULE, HAIR, ROSE, BLUSH,
                   DOCX_SERIF, DOCX_SANS)

SIZES = {"letter": ("8.5in", "11in"), "a4": ("210mm", "297mm")}


def esc(s):
    return _h.escape(str(s), quote=True)


def inline(s):
    s = esc(s)
    s = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s)
    s = re.sub(r"(?<![\w/])_(.+?)_(?![\w/])", r"<i>\1</i>", s)
    return s


# ───────────────────────────── HTML ─────────────────────────────

CSS = """
*{box-sizing:border-box}
html,body{margin:0;padding:0;background:#fff}
body{color:var(--ink);font-family:var(--sans);font-weight:300;font-size:8.8pt;line-height:1.45;
  -webkit-print-color-adjust:exact;print-color-adjust:exact;font-kerning:normal}
.page{width:var(--pw);height:var(--ph);padding:.5in .6in .62in;position:relative;overflow:hidden;
  break-after:page;display:flex;flex-direction:column}
.page:last-child{break-after:auto}
.top{display:flex;justify-content:space-between;align-items:flex-end;gap:18pt;margin:0 0 12pt;
  font-size:7.2pt;letter-spacing:.14em;text-transform:uppercase;color:var(--muted);font-weight:400}
.top .studio{display:flex;align-items:flex-end;gap:6pt;flex:0 1 3.3in}
.top .studio .ln{flex:1}
.top .mark{font-family:var(--serif);font-size:11pt;letter-spacing:.2em;color:var(--ink);font-weight:500}
.foot{position:absolute;left:.6in;right:.6in;bottom:.32in;display:flex;justify-content:space-between;
  font-size:6.8pt;letter-spacing:.12em;text-transform:uppercase;color:var(--muted);font-weight:400}
.title{margin:0 0 4pt}
.title h1{font-family:var(--serif);font-weight:500;font-size:25pt;line-height:1.02;letter-spacing:-.005em;margin:0;text-wrap:balance}
.title h1 i{color:var(--rose);font-weight:500}
.title p{margin:5pt 0 0;font-size:9pt;color:var(--muted);max-width:5.6in}
.h{display:flex;align-items:center;gap:8pt;margin:13pt 0 7pt;font-size:7.4pt;font-weight:500;
  letter-spacing:.16em;text-transform:uppercase;color:var(--rose);break-after:avoid}
.h::after{content:"";flex:1;border-top:.6pt solid var(--hair)}
p.p{margin:0 0 6pt;max-width:6.9in}
p.small{margin:4pt 0 6pt;font-size:7.3pt;line-height:1.4;color:var(--muted)}
ul.list,ol.steps{margin:0 0 7pt;padding-left:13pt}
ul.list li,ol.steps li{margin:0 0 2.5pt;padding-left:2pt}
ul.list li::marker{color:var(--rose)}
ol.steps li::marker{color:var(--rose);font-weight:500}
b{font-weight:500}
i{font-style:italic}
.row{display:flex;gap:14pt;margin:0 0 6pt}
.fld{display:flex;align-items:flex-end;gap:6pt;min-width:0}
.lb{font-size:8pt;color:var(--muted);white-space:nowrap;font-weight:400;line-height:1.1}
.ln{display:block;flex:1;min-width:.35in;height:15pt;border-bottom:.75pt solid var(--rule)}
.lines{margin:0 0 6pt}
.lines .lb{display:block;margin:0 0 1pt}
.lines .ln{height:17pt}
.checks{display:flex;gap:10pt;margin:0 0 6pt;align-items:flex-start}
.checks>.lb{padding-top:1pt;flex:0 0 1.28in;white-space:normal}
.cgrid{flex:1;display:grid;grid-template-columns:repeat(var(--c),minmax(0,1fr));gap:3.5pt 10pt;font-size:8.5pt}
.cgrid span.o{display:flex;align-items:center;gap:5pt;line-height:1.2}
.cb{display:inline-block;flex:0 0 auto;width:8.4pt;height:8.4pt;border:.75pt solid #968690;border-radius:1.8pt;background:#fff}
.yn{display:grid;grid-template-columns:minmax(0,1fr) auto auto 2.15in;column-gap:10pt;margin:0 0 6pt;font-size:8.5pt}
.yn>div{border-bottom:.5pt solid var(--hair);padding:3.3pt 0 2.9pt;display:flex;align-items:center;gap:4pt;line-height:1.25}
.yn .hd{font-size:6.8pt;letter-spacing:.12em;text-transform:uppercase;color:var(--muted);font-weight:400;border-bottom:.6pt solid var(--rose)}
.yn .d .ln{height:12pt;border-bottom:none}
.ini{display:grid;grid-template-columns:minmax(0,1fr) 1.05in;gap:14pt;align-items:end;margin:0 0 6.5pt}
.ini p{margin:0}
.ini .fld .lb{font-size:7pt;letter-spacing:.1em;text-transform:uppercase}
.sign{margin:10pt 0 4pt}
.sign .ln{height:24pt}
table.t{width:100%;border-collapse:collapse;margin:0 0 7pt;font-size:8pt;table-layout:fixed}
table.t th{text-align:left;font-weight:400;font-size:6.8pt;letter-spacing:.11em;text-transform:uppercase;color:var(--muted);
  border-bottom:.75pt solid var(--rose);padding:0 5pt 3.5pt 0;vertical-align:bottom;line-height:1.2}
table.t td{border-bottom:.5pt solid var(--hair);padding:3pt 5pt 3pt 0;vertical-align:top;height:var(--rh)}
table.t td.c{text-align:center;padding-right:0}
table.t td.c .cb{vertical-align:middle}
table.t.grid td,table.t.grid th{border-right:.5pt solid var(--hair);padding-left:5pt}
table.t.grid td:last-child,table.t.grid th:last-child{border-right:none}
.svg{margin:2pt 0 6pt;display:flex;flex-direction:column;align-items:center}
.svg svg{height:100%;width:auto;display:block}
.svg .cap{font-size:7pt;color:var(--muted);margin-top:3pt;text-align:center}
.cols{display:grid;gap:22pt;margin:0}
.cols>div{min-width:0}
.script{position:relative;margin:0 0 9pt;padding:0 0 0 16pt;break-inside:avoid}
.script .when{font-size:6.8pt;letter-spacing:.13em;text-transform:uppercase;color:var(--muted);font-weight:400;margin:0 0 1pt}
.script .say{font-family:var(--serif);font-style:italic;font-size:12pt;line-height:1.3;color:var(--ink);margin:0}
.script::before{content:"\\201C";position:absolute;left:0;top:5pt;font-family:var(--serif);font-size:26pt;line-height:1;color:var(--rose)}
.callout{background:var(--blush);border-radius:4pt;padding:9pt 12pt 8pt;margin:4pt 0 9pt;break-inside:avoid}
.callout .ct{font-family:var(--serif);font-size:12.5pt;font-weight:600;margin:0 0 2pt}
.callout p{margin:0}
.box{border:.75pt solid var(--rule);border-radius:3pt;padding:5pt 7pt;margin:0 0 7pt;position:relative}
.box .lb{font-size:6.8pt;letter-spacing:.12em;text-transform:uppercase}
.box .area{position:absolute;left:4pt;right:4pt;bottom:4pt;top:16pt}
.kv{display:grid;grid-template-columns:1.35in minmax(0,1fr);gap:5pt 14pt;margin:0 0 8pt}
.kv dt{font-weight:500;color:var(--ink)}
.kv dd{margin:0}
.mast{text-align:center;margin:.1in 0 .16in}
.mname{font-family:var(--serif);font-weight:500;font-size:30pt;letter-spacing:.2em;text-transform:uppercase;line-height:1}
.mtag{margin-top:7pt;font-size:7.4pt;letter-spacing:.3em;text-transform:uppercase;color:var(--rose);font-weight:400}
.menu{display:grid;gap:9pt;margin:0 0 6pt}
.mi p{margin:1.5pt 0 0;color:var(--muted);font-size:8.5pt;max-width:6.3in}
.mh{display:flex;align-items:baseline;gap:7pt}
.mn{font-family:var(--serif);font-size:14.5pt;font-weight:600;line-height:1.1}
.mt{font-size:6.8pt;letter-spacing:.14em;text-transform:uppercase;color:var(--muted);font-weight:400;white-space:nowrap}
.md{flex:1;border-bottom:.75pt dotted #b9a9ae;transform:translateY(-2.5pt);min-width:.3in}
.mp{font-weight:400;font-size:9.5pt;font-variant-numeric:tabular-nums;white-space:nowrap}
.prices{display:grid;grid-template-columns:repeat(var(--c),minmax(0,1fr));grid-template-rows:repeat(var(--r),auto);grid-auto-flow:column;gap:4.5pt 22pt;margin:0 0 6pt}
.pr{display:flex;align-items:baseline;gap:6pt;font-size:9pt;font-weight:400}
"""


def _ln(name, kind="text", extra=""):
    attr = ' data-f="%s" data-t="%s"' % (esc(name), kind) if name else ""
    return '<span class="ln"%s%s></span>' % (attr, extra)


def _cb(name):
    return '<span class="cb" data-f="%s" data-t="check"></span>' % esc(name) if name else '<span class="cb"></span>'


def slug(s):
    return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")[:28]


def html_block(b):
    k = b[0]
    if k == "title":
        sub = '<p>%s</p>' % inline(b[2]) if len(b) > 2 and b[2] else ""
        t = inline(b[1])
        return '<div class="title"><h1>%s</h1>%s</div>' % (t, sub)
    if k == "h":
        return '<div class="h">%s</div>' % inline(b[1])
    if k == "p":
        return '<p class="p">%s</p>' % inline(b[1])
    if k == "small":
        return '<p class="small">%s</p>' % inline(b[1])
    if k == "list":
        return '<ul class="list">%s</ul>' % "".join("<li>%s</li>" % inline(x) for x in b[1])
    if k == "steps":
        return '<ol class="steps">%s</ol>' % "".join("<li>%s</li>" % inline(x) for x in b[1])
    if k == "fields":
        cells = []
        for label, w, name in b[1]:
            lb = '<span class="lb">%s</span>' % inline(label) if label else ""
            cells.append('<div class="fld" style="flex:%s">%s%s</div>' % (w, lb, _ln(name)))
        return '<div class="row">%s</div>' % "".join(cells)
    if k == "lines":
        n, name = b[1], b[2]
        label = b[3] if len(b) > 3 else ""
        lb = '<span class="lb">%s</span>' % inline(label) if label else ""
        return ('<div class="lines" data-f="%s" data-t="multi">%s%s</div>'
                % (esc(name), lb, "".join('<span class="ln"></span>' for _ in range(n))))
    if k == "checks":
        label, opts, name, cols = b[1], b[2], b[3], b[4]
        lb = '<span class="lb">%s</span>' % inline(label) if label else ""
        items = []
        for o in opts:
            if o.endswith("___"):
                items.append('<span class="o">%s<span>%s</span>%s</span>' % (
                    _cb("%s_%s" % (name, slug(o))), inline(o[:-3].strip()), _ln("%s_%s_text" % (name, slug(o)))))
            else:
                items.append('<span class="o">%s<span>%s</span></span>' % (_cb("%s_%s" % (name, slug(o))), inline(o)))
        return '<div class="checks">%s<div class="cgrid" style="--c:%d">%s</div></div>' % (lb, cols, "".join(items))
    if k == "yn":
        out = ['<div class="hd"></div><div class="hd">Yes</div><div class="hd">No</div><div class="hd">If yes, details</div>']
        for q, name in b[1]:
            out.append('<div>%s</div><div>%s</div><div>%s</div><div class="d">%s</div>' % (
                inline(q), _cb(name + "_yes"), _cb(name + "_no"), _ln(name + "_details")))
        return '<div class="yn">%s</div>' % "".join(out)
    if k == "initial":
        return ('<div class="ini"><p>%s</p><div class="fld"><span class="lb">Initials</span>%s</div></div>'
                % (inline(b[1]), _ln(b[2])))
    if k == "sign":
        cells = []
        for label, w, name in b[1]:
            cells.append('<div class="fld" style="flex:%s"><span class="lb">%s</span>%s</div>' % (w, inline(label), _ln(name)))
        return '<div class="row sign">%s</div>' % "".join(cells)
    if k == "table":
        heads, widths, rows, name = b[1], b[2], b[3], b[4]
        rh = b[5] if len(b) > 5 else 17
        cls = "t grid" if len(b) > 6 and b[6] == "grid" else "t"
        col = "".join('<col style="width:%s%%">' % w for w in widths)
        th = "".join("<th>%s</th>" % inline(x) for x in heads)
        body = []
        if isinstance(rows, int):
            for r in range(rows):
                tds = []
                for c, hname in enumerate(heads):
                    if hname == "✓":
                        tds.append('<td class="c">%s</td>' % _cb("%s_r%d_c%d" % (name, r + 1, c + 1)))
                    else:
                        tds.append('<td data-f="%s_r%d_c%d" data-t="cell"></td>' % (esc(name), r + 1, c + 1))
                body.append("<tr>%s</tr>" % "".join(tds))
        else:
            for r, row in enumerate(rows):
                tds = []
                for c, v in enumerate(row):
                    if v is None:
                        tds.append('<td data-f="%s_r%d_c%d" data-t="cell"></td>' % (esc(name), r + 1, c + 1))
                    elif v == "[]":
                        tds.append('<td class="c">%s</td>' % _cb("%s_r%d_c%d" % (name, r + 1, c + 1)))
                    else:
                        tds.append("<td>%s</td>" % inline(v))
                body.append("<tr>%s</tr>" % "".join(tds))
        return '<table class="%s" style="--rh:%spt"><colgroup>%s</colgroup><thead><tr>%s</tr></thead><tbody>%s</tbody></table>' % (
            cls, rh, col, th, "".join(body))
    if k == "svg":
        cap = '<div class="cap">%s</div>' % inline(b[3]) if len(b) > 3 and b[3] else ""
        return '<div class="svg"><div style="height:%sin">%s</div>%s</div>' % (b[2], b[1], cap)
    if k == "cols":
        left, right, share = b[1], b[2], b[3] if len(b) > 3 else 0.5
        return ('<div class="cols" style="grid-template-columns:minmax(0,%sfr) minmax(0,%sfr)"><div>%s</div><div>%s</div></div>'
                % (share, 1 - share, "".join(html_block(x) for x in left), "".join(html_block(x) for x in right)))
    if k == "script":
        return '<div class="script"><div class="when">%s</div><p class="say">%s</p></div>' % (inline(b[1]), inline(b[2]))
    if k == "callout":
        return '<div class="callout"><div class="ct">%s</div><p>%s</p></div>' % (inline(b[1]), inline(b[2]))
    if k == "box":
        return ('<div class="box" style="height:%sin"><span class="lb">%s</span><div class="area" data-f="%s" data-t="multi"></div></div>'
                % (b[2], inline(b[1]), esc(b[3])))
    if k == "kv":
        return '<dl class="kv">%s</dl>' % "".join("<dt>%s</dt><dd>%s</dd>" % (inline(a), inline(c)) for a, c in b[1])
    if k == "space":
        return '<div style="height:%sin"></div>' % b[1]
    if k == "masthead":
        return '<div class="mast"><div class="mname">%s</div><div class="mtag">%s</div></div>' % (inline(b[1]), inline(b[2]))
    if k == "menu":
        items = []
        for name, mins, price, desc in b[1]:
            items.append('<div class="mi"><div class="mh"><span class="mn">%s</span><span class="mt">%s</span>'
                         '<span class="md"></span><span class="mp">%s</span></div><p>%s</p></div>'
                         % (inline(name), inline(mins), inline(price), inline(desc)))
        return '<div class="menu">%s</div>' % "".join(items)
    if k == "prices":
        cols = b[2] if len(b) > 2 else 2
        items = "".join('<div class="pr"><span>%s</span><span class="md"></span><span class="mp">%s</span></div>'
                        % (inline(a), inline(c)) for a, c in b[1])
        return '<div class="prices" style="--c:%d;--r:%d">%s</div>' % (cols, (len(b[1]) + cols - 1) // cols, items)
    raise ValueError("unknown block " + k)


def page_html(doc, page, i, n):
    head = doc.get("header")
    top = ""
    if head == "studio":
        top = ('<div class="top"><div class="studio fld"><span class="lb" style="font-size:7.2pt;letter-spacing:.14em">Studio</span>%s</div><span>%s</span></div>'
               % (_ln("studio_p%d" % (i + 1)), esc(page["name"])))
    elif head == "brand":
        top = '<div class="top"><span class="mark">LUMEVINA</span><span>%s</span></div>' % esc(page.get("kicker", doc["title"]))
    foot_l = doc.get("foot", doc["title"])
    foot = '<div class="foot"><span>%s</span><span>%d / %d</span></div>' % (esc(foot_l), i + 1, n)
    body = "".join(html_block(b) for b in page["blocks"])
    return '<section class="page %s">%s%s%s</section>' % (page.get("cls", ""), top, body, foot)


def html(doc, size="letter", extra_css=""):
    pw, ph = SIZES[size]
    pages = doc["pages"]
    body = "".join(page_html(doc, p, i, len(pages)) for i, p in enumerate(pages))
    return ('<!doctype html><html lang="en"><head><meta charset="utf-8"><title>%s</title><style>%s\n%s\n'
            ':root{--pw:%s;--ph:%s}@page{size:%s %s;margin:0}%s\n%s</style></head><body>%s</body></html>'
            % (esc(doc["title"]), font_css(), TOKENS, pw, ph, pw, ph, CSS, extra_css, body))


# ───────────────────────────── Word ─────────────────────────────

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor


def rgb(hexs):
    hexs = hexs.lstrip("#")
    return RGBColor(int(hexs[0:2], 16), int(hexs[2:4], 16), int(hexs[4:6], 16))


def _font(run, name=DOCX_SANS, size=9, color=INK, bold=False, italic=False, caps=False, spacing=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.font.bold = bold
    run.font.italic = italic
    run.font.all_caps = caps
    if spacing is not None:
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(int(spacing * 20)))
        rpr.append(sp)
    return run


def _para(container, first=None):
    """A new paragraph; in a fresh table cell, reuse the empty one it starts with."""
    if first is not None:
        return first
    return container.add_paragraph()


def _fmt(p, before=0, after=4, line=None, align=None, keep=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line
    if align:
        pf.alignment = align
    if keep:
        pf.keep_with_next = True
    return p


def _rich(p, text, size=9, color=INK, name=DOCX_SANS, italic=False):
    parts = re.split(r"(\*\*.+?\*\*|(?<![\w/])_.+?_(?![\w/]))", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            _font(p.add_run(part[2:-2]), name, size, color, bold=True, italic=italic)
        elif part.startswith("_") and part.endswith("_") and len(part) > 2:
            _font(p.add_run(part[1:-1]), name, size, color, italic=True)
        else:
            _font(p.add_run(part), name, size, color, italic=italic)


def _border(el, edges, color=RULE, sz=6, tag="w:tcBorders", pr=None):
    pr = pr if pr is not None else el
    b = pr.find(qn(tag))
    if b is None:
        b = OxmlElement(tag)
        pr.append(b)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        if edge in edges:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:color"), color.lstrip("#"))
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)


def _cell_border(cell, edges, color=RULE, sz=6):
    tcpr = cell._tc.get_or_add_tcPr()
    _border(None, edges, color, sz, "w:tcBorders", tcpr)


def _table_plain(t):
    tblpr = t._tbl.tblPr
    _border(None, (), tag="w:tblBorders", pr=tblpr)
    cm = OxmlElement("w:tblCellMar")
    for edge, v in (("left", 0), ("right", 90), ("top", 0), ("bottom", 0)):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        cm.append(e)
    tblpr.append(cm)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT


def _row_height(row, pts, exact=False):
    trpr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(pts * 20)))
    h.set(qn("w:hRule"), "exact" if exact else "atLeast")
    trpr.append(h)


def _widths(t, inches):
    """Fixed column widths that Word and LibreOffice both honor: grid, table width and every cell."""
    t.autofit = False
    tbl = t._tbl
    grid = tbl.tblGrid
    cols = grid.findall(qn("w:gridCol"))
    for gc, w in zip(cols, inches):
        gc.set(qn("w:w"), str(int(w * 1440)))
    tblpr = tbl.tblPr
    tw = tblpr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW")
        tblpr.append(tw)
    tw.set(qn("w:w"), str(int(sum(inches) * 1440)))
    tw.set(qn("w:type"), "dxa")
    lay = tblpr.find(qn("w:tblLayout"))
    if lay is None:
        lay = OxmlElement("w:tblLayout")
        tblpr.append(lay)
    lay.set(qn("w:type"), "fixed")
    for row in t.rows:
        for c, w in zip(row.cells, inches):
            c.width = Inches(w)


def _gap(c, pts):
    """The paragraph Word needs between tables, kept as short as the space wanted."""
    p = c.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = Pt(max(pts, 1))
    from docx.enum.text import WD_LINE_SPACING
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    _font(p.add_run(""), size=1)
    return p


def _cell_va(cell, v="bottom"):
    tcpr = cell._tc.get_or_add_tcPr()
    va = OxmlElement("w:vAlign")
    va.set(qn("w:val"), v)
    tcpr.append(va)


def _shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), fill.lstrip("#"))
    tcpr.append(sh)


BOX = "☐"   # ☐


ORDER = {
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid "
           "vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs "
           "em lang eastAsianLayout specVanish oMath",
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs "
           "suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi "
           "adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection "
           "textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing "
             "tblInd tblBorders shd tblLayout tblCellMar tblLook",
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark",
    "trPr": "cnfStyle divId gridBefore gridAfter wBefore wAfter cantSplit trHeight tblHeader tblCellSpacing jc hidden",
}
ORDER = {k: {n: i for i, n in enumerate(v.split())} for k, v in ORDER.items()}


def _normalize(root):
    """Put property children in the order the Word schema requires (Word refuses files that don't)."""
    for tag, rank in ORDER.items():
        for pr in root.iter(qn("w:" + tag)):
            kids = list(pr)
            ordered = sorted(kids, key=lambda e: rank.get(e.tag.split("}")[1], 999))
            if ordered != kids:
                for e in kids:
                    pr.remove(e)
                for e in ordered:
                    pr.append(e)


class Word:
    def __init__(self, doc, size="letter", width_in=None):
        self.d = Document()
        self.doc = doc
        s = self.d.sections[0]
        margin = 0.6
        if size == "letter":
            s.page_width, s.page_height = Inches(8.5), Inches(11)
        elif size == "a4":
            s.page_width, s.page_height = Inches(8.27), Inches(11.69)
        else:                                   # (width, height, margin) in inches
            s.page_width, s.page_height, margin = Inches(size[0]), Inches(size[1]), size[2]
        s.left_margin = s.right_margin = Inches(margin)
        s.top_margin = Inches(min(0.5, margin))
        s.bottom_margin = Inches(min(0.55, margin + 0.1))
        s.footer_distance = Inches(0.3)
        self.W = (s.page_width - s.left_margin - s.right_margin) / 914400.0
        st = self.d.styles["Normal"]
        st.font.name = DOCX_SANS
        st.font.size = Pt(9)
        st.paragraph_format.space_after = Pt(4)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.insert(0, rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), DOCX_SANS)
        fp = s.footer.paragraphs[0]
        _font(fp.add_run(doc.get("foot", doc["title"]).upper()), size=6.8, color=MUTED, spacing=1)

    # each block writer takes a container (document or table cell) and width in inches
    def page(self, page, first):
        c = self.d
        if not first:
            p = c.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)
            _fmt(p, 0, 0)
        head = self.doc.get("header")
        if head == "studio":
            t = c.add_table(rows=1, cols=3)
            _table_plain(t)
            a, b, d = t.rows[0].cells
            _font(a.paragraphs[0].add_run("STUDIO"), size=7, color=MUTED, spacing=1.5)
            _cell_border(b, ("bottom",))
            _font(d.paragraphs[0].add_run(page["name"].upper()), size=7, color=MUTED, spacing=1.5)
            d.paragraphs[0].alignment = 2
            _widths(t, [0.6, 2.6, self.W - 3.2])
            _gap(c, 6)
        elif head == "brand":
            p = c.add_paragraph()
            _font(p.add_run("LUMEVINA"), DOCX_SERIF, 11, INK, bold=True, spacing=3)
            _font(p.add_run("\t" + page.get("kicker", self.doc["title"]).upper()), size=7, color=MUTED, spacing=1.5)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(self.W), 2)
            _fmt(p, 0, 8)
        for b in page["blocks"]:
            self.block(c, b, self.W)

    def block(self, c, b, W):
        k = b[0]
        if k == "title":
            p = _fmt(c.add_paragraph(), 0, 2, keep=True)
            _rich(p, b[1], 24, INK, DOCX_SERIF)
            if len(b) > 2 and b[2]:
                p = _fmt(c.add_paragraph(), 0, 6)
                _rich(p, b[2], 9, MUTED)
        elif k == "h":
            p = _fmt(c.add_paragraph(), 10, 5, keep=True)
            _font(p.add_run(re.sub(r"[*_]", "", b[1]).upper()), size=7.4, color=ROSE, bold=True, spacing=1.8)
            pbdr = OxmlElement("w:pBdr")
            bot = OxmlElement("w:bottom")
            for a, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "2"), ("w:color", HAIR.lstrip("#"))):
                bot.set(qn(a), v)
            pbdr.append(bot)
            p._p.get_or_add_pPr().append(pbdr)
        elif k == "p":
            _rich(_fmt(c.add_paragraph(), 0, 5, 1.15), b[1], 9)
        elif k == "small":
            _rich(_fmt(c.add_paragraph(), 2, 5), b[1], 7.3, MUTED)
        elif k in ("list", "steps"):
            for n, x in enumerate(b[1]):
                p = _fmt(c.add_paragraph(), 0, 2.5, 1.1)
                p.paragraph_format.left_indent = Inches(0.2)
                p.paragraph_format.first_line_indent = Inches(-0.16)
                _font(p.add_run(("%d.\t" % (n + 1)) if k == "steps" else "•\t"), size=9, color=ROSE, bold=k == "steps")
                p.paragraph_format.tab_stops.add_tab_stop(Inches(0.2))
                _rich(p, x, 9)
        elif k in ("fields", "sign"):
            cells = b[1]
            tot = float(sum(w for _, w, _ in cells))
            t = c.add_table(rows=1, cols=len(cells) * 2)
            _table_plain(t)
            ws = []
            for i, (label, w, name) in enumerate(cells):
                lc, vc = t.rows[0].cells[2 * i], t.rows[0].cells[2 * i + 1]
                lw = 0.2 + 0.062 * len(label) if label else 0.02
                share = W * w / tot
                ws += [lw, max(share - lw, 0.3)]
                if label:
                    _font(lc.paragraphs[0].add_run(label), size=8, color=MUTED)
                _cell_border(vc, ("bottom",))
                _cell_va(lc)
            _widths(t, ws)
            _row_height(t.rows[0], 24 if k == "sign" else 15)
            _gap(c, 4)
        elif k == "lines":
            if len(b) > 3 and b[3]:
                _font(_fmt(c.add_paragraph(), 0, 0).add_run(b[3]), size=8, color=MUTED)
            t = c.add_table(rows=b[1], cols=1)
            _table_plain(t)
            for r in t.rows:
                _cell_border(r.cells[0], ("bottom",))
                _row_height(r, 17)
            _widths(t, [W])
            _gap(c, 4)
        elif k == "checks":
            label, opts, name, cols = b[1], b[2], b[3], b[4]
            if label:
                _font(_fmt(c.add_paragraph(), 0, 1, keep=True).add_run(label), size=8, color=MUTED)
            rows = (len(opts) + cols - 1) // cols
            t = c.add_table(rows=rows, cols=cols)
            _table_plain(t)
            for i, o in enumerate(opts):
                cell = t.rows[i // cols].cells[i % cols]
                p = cell.paragraphs[0]
                _font(p.add_run(BOX + "  "), size=10, color=MUTED)
                _font(p.add_run(o.replace("___", " ________")), size=8.5)
                _fmt(p, 0, 1)
            _widths(t, [W / cols] * cols)
            _gap(c, 4)
        elif k == "yn":
            t = c.add_table(rows=len(b[1]) + 1, cols=4)
            _table_plain(t)
            for i, h in enumerate(("", "Yes", "No", "If yes, details")):
                cell = t.rows[0].cells[i]
                _font(cell.paragraphs[0].add_run(h.upper()), size=6.8, color=MUTED, spacing=1)
                _cell_border(cell, ("bottom",), ROSE, 6)
            for r, (q, name) in enumerate(b[1]):
                row = t.rows[r + 1]
                _rich(row.cells[0].paragraphs[0], q, 8.5)
                for j in (1, 2):
                    _font(row.cells[j].paragraphs[0].add_run(BOX), size=10, color=MUTED)
                for cell in row.cells:
                    _cell_border(cell, ("bottom",), HAIR, 4)
                    _fmt(cell.paragraphs[0], 2, 2)
            _widths(t, [W - 2.95, 0.4, 0.4, 2.15])
            _gap(c, 4)
        elif k == "initial":
            t = c.add_table(rows=1, cols=3)
            _table_plain(t)
            a, lb, v = t.rows[0].cells
            _rich(a.paragraphs[0], b[1], 8.8)
            _font(lb.paragraphs[0].add_run("INITIALS"), size=7, color=MUTED, spacing=1)
            _cell_border(v, ("bottom",))
            _cell_va(lb)
            _widths(t, [W - 1.4, 0.72, 0.68])
            _gap(c, 2)
        elif k == "table":
            heads, widths, rows, name = b[1], b[2], b[3], b[4]
            rh = b[5] if len(b) > 5 else 17
            n = rows if isinstance(rows, int) else len(rows)
            t = c.add_table(rows=n + 1, cols=len(heads))
            _table_plain(t)
            for i, h in enumerate(heads):
                cell = t.rows[0].cells[i]
                _font(cell.paragraphs[0].add_run(h.upper() if h != "✓" else "✓"), size=6.8, color=MUTED, spacing=1)
                _cell_border(cell, ("bottom",), ROSE, 6)
            for r in range(n):
                row = t.rows[r + 1]
                _row_height(row, rh)
                for i, cell in enumerate(row.cells):
                    _cell_border(cell, ("bottom",), HAIR, 4)
                    v = None if isinstance(rows, int) else rows[r][i]
                    if isinstance(rows, int) and heads[i] == "✓" or v == "[]":
                        _font(cell.paragraphs[0].add_run(BOX), size=10, color=MUTED)
                        cell.paragraphs[0].alignment = 1
                    elif v:
                        _rich(cell.paragraphs[0], v, 8)
                    _fmt(cell.paragraphs[0], 2, 1)
            _widths(t, [W * w / 100.0 for w in widths])
            _gap(c, 6)
        elif k == "svg":
            png = self.doc.get("_png", {}).get(id(b))
            if png:
                p = _fmt(c.add_paragraph(), 2, 2, align=1)
                p.add_run().add_picture(png, height=Inches(b[2]))
            if len(b) > 3 and b[3]:
                _rich(_fmt(c.add_paragraph(), 0, 6, align=1), b[3], 7, MUTED)
        elif k == "cols":
            left, right, share = b[1], b[2], b[3] if len(b) > 3 else 0.5
            t = c.add_table(rows=1, cols=2)
            _table_plain(t)
            lw, rw = (W - 0.3) * share, (W - 0.3) * (1 - share)
            for cell, blocks, w in ((t.rows[0].cells[0], left, lw), (t.rows[0].cells[1], right, rw)):
                for x in blocks:
                    self.block(cell, x, w)
                # a cell must end in a paragraph; drop the empty one it started with if something came first
                first = cell.paragraphs[0]
                if not first.text and len(cell._tc) > 2 and first._p is cell._tc[1]:
                    cell._tc.remove(first._p)
            _widths(t, [lw + 0.15, rw + 0.15])
            _gap(c, 4)
        elif k == "script":
            p = _fmt(c.add_paragraph(), 3, 0, keep=True)
            _font(p.add_run(b[1].upper()), size=6.8, color=MUTED, spacing=1)
            p.paragraph_format.left_indent = Inches(0.22)
            p = _fmt(c.add_paragraph(), 0, 7, 1.1)
            p.paragraph_format.left_indent = Inches(0.22)
            _font(p.add_run("“"), DOCX_SERIF, 13, ROSE)
            _rich(p, b[2], 12, INK, DOCX_SERIF, italic=True)
            _font(p.add_run("”"), DOCX_SERIF, 13, ROSE)
        elif k == "callout":
            t = c.add_table(rows=1, cols=1)
            _table_plain(t)
            cell = t.rows[0].cells[0]
            _shade(cell, BLUSH)
            p = cell.paragraphs[0]
            _fmt(p, 5, 1)
            p.paragraph_format.left_indent = Inches(0.1)
            _rich(p, b[1], 12.5, INK, DOCX_SERIF)
            p = _fmt(cell.add_paragraph(), 0, 6)
            p.paragraph_format.left_indent = Inches(0.1)
            p.paragraph_format.right_indent = Inches(0.1)
            _rich(p, b[2], 9)
            _widths(t, [W])
            _gap(c, 6)
        elif k == "box":
            t = c.add_table(rows=1, cols=1)
            _table_plain(t)
            cell = t.rows[0].cells[0]
            _cell_border(cell, ("top", "left", "bottom", "right"))
            p = cell.paragraphs[0]
            p.paragraph_format.left_indent = Inches(0.06)
            _font(p.add_run(b[1].upper()), size=6.8, color=MUTED, spacing=1)
            _row_height(t.rows[0], b[2] * 72)
            _widths(t, [W])
            _gap(c, 6)
        elif k == "kv":
            t = c.add_table(rows=len(b[1]), cols=2)
            _table_plain(t)
            for i, (a, d) in enumerate(b[1]):
                _rich(t.rows[i].cells[0].paragraphs[0], "**%s**" % a, 9)
                _rich(t.rows[i].cells[1].paragraphs[0], d, 9)
                for cell in t.rows[i].cells:
                    _fmt(cell.paragraphs[0], 0, 4, 1.1)
            _widths(t, [1.4, W - 1.4])
            _gap(c, 6)
        elif k == "space":
            _fmt(c.add_paragraph(), 0, b[1] * 72)
        elif k == "masthead":
            p = _fmt(c.add_paragraph(), 6, 2, align=1)
            _font(p.add_run(re.sub(r"[*_]", "", b[1]).upper()), DOCX_SERIF, 28, INK, spacing=5)
            p = _fmt(c.add_paragraph(), 0, 12, align=1)
            _font(p.add_run(re.sub(r"[*_]", "", b[2]).upper()), size=7.4, color=ROSE, spacing=3)
        elif k == "menu":
            for name, mins, price, desc in b[1]:
                p = _fmt(c.add_paragraph(), 5, 0, keep=True)
                p.paragraph_format.tab_stops.add_tab_stop(Inches(W), 2, 1)   # right-aligned, dotted leader
                _rich(p, name, 14, INK, DOCX_SERIF)
                _font(p.add_run("   " + mins.upper()), size=6.8, color=MUTED, spacing=1)
                _font(p.add_run("\t" + price), size=9.5)
                _rich(_fmt(c.add_paragraph(), 1, 4, 1.1), desc, 8.5, MUTED)
        elif k == "prices":
            cols = b[2] if len(b) > 2 else 2
            rows = (len(b[1]) + cols - 1) // cols
            t = c.add_table(rows=rows, cols=cols)
            _table_plain(t)
            cw = W / cols
            for i, (a, pr) in enumerate(b[1]):
                cell = t.rows[i % rows].cells[i // rows]
                p = cell.paragraphs[0]
                p.paragraph_format.tab_stops.add_tab_stop(Inches(cw - 0.3), 2, 1)
                _rich(p, a, 9)
                _font(p.add_run("\t" + pr), size=9)
                _fmt(p, 1, 2)
            _widths(t, [cw] * cols)
            _gap(c, 6)
        else:
            raise ValueError(k)

    def save(self, path):
        for i, page in enumerate(self.doc["pages"]):
            self.page(page, i == 0)
        self.d.core_properties.title = self.doc["title"]
        self.d.core_properties.author = "Lumevina Studio"
        _normalize(self.d.element)
        for sec in self.d.sections:
            _normalize(sec.footer._element)
        _normalize(self.d.styles.element)
        self.d.save(path)


def docx(doc, path, size="letter"):
    Word(doc, size).save(path)
