"""05 · Membership Pricing Calculator: design tiers, see what each member is worth, and how many you need."""
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill

import sheets as S

SLUG = "05-membership-pricing-calculator"
NAME = "Membership Pricing Calculator"


def start_sheet(ws, name, steps, extra=None):
    S.title(ws, name, "Made by a licensed esthetician. Works in Excel, Google Sheets and Numbers.")
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 90
    S.head(ws, 4, "How to use it", 2)
    for i, s in enumerate(steps):
        ws.cell(row=5 + i, column=1, value="Step %d" % (i + 1)).font = S.BOLD
        c = ws.cell(row=5 + i, column=2, value=s)
        c.font = S.LABEL
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[5 + i].height = 30
    r = 6 + len(steps)
    S.head(ws, r, "Which cells to change", 2)
    S.legend(ws, r + 1)
    r += 5
    for heading, text in (extra or []):
        S.head(ws, r, heading, 2)
        c = ws.cell(row=r + 1, column=1, value=text)
        c.font = S.LABEL
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=2)
        ws.row_dimensions[r + 1].height = 44
        r += 3
    return r


def calculator(ws):
    S.title(ws, "Membership Calculator", "Every blue number is an example. Replace it with yours.")
    for col, w in (("A", 40), ("B", 14), ("C", 14), ("D", 14), ("E", 62)):
        ws.column_dimensions[col].width = w

    S.head(ws, 4, "Your monthly costs", 5)
    costs = [("Rent and utilities", 1400, "Suite or room rent, plus power and water if you pay them."),
             ("Insurance and license", 65, "Liability insurance and license renewals, spread monthly."),
             ("Booking software and website", 120, "Booking app, website, email, phone."),
             ("Marketing", 150, "Ads, printing, promotions."),
             ("Laundry and small supplies", 60, "Linens and disposables not tied to one service."),
             ("Other", 0, "Anything else you pay every month.")]
    for i, (lab, v, n) in enumerate(costs):
        r = 5 + i
        S.label(ws, "A%d" % r, lab)
        S.inp(ws, "B%d" % r, v, S.USD, key=(i == 0))
        S.note(ws, "E%d" % r, n)
    S.label(ws, "A11", "Total monthly costs", bold=True)
    S.calc(ws, "B11", "=SUM(B5:B10)", S.USD, bold=True)

    S.head(ws, 13, "Card fees", 5)
    S.label(ws, "A14", "Percent of each charge")
    S.inp(ws, "B14", 0.029, S.PCT1)
    S.note(ws, "E14", "A typical online card rate (Stripe, Square). Check your own processor.")
    S.label(ws, "A15", "Fixed fee per charge")
    S.inp(ws, "B15", 0.30, S.USD2)

    S.head(ws, 17, "Your membership tiers", 5)
    for col, t in (("B", "Tier 1"), ("C", "Tier 2"), ("D", "Tier 3")):
        c = ws["%s17" % col]
        c.value = t.upper()
        c.font = S.HEAD
        c.alignment = Alignment(horizontal="right")
    tiers = [
        ("Name", ["Glow", "Glow Plus", "Ageless"], None, "What clients will see.", False),
        ("Monthly price", [119, 159, 209], S.USD, "What the member pays each month.", True),
        ("Included service: your single price", [140, 185, 245], S.USD, "What the included facial costs without a membership.", False),
        ("Minutes per included service", [60, 75, 90], S.NUM, "Treatment time, including turnover.", False),
        ("Product used per service (your cost)", [12, 16, 22], S.USD, "Backbar product used in the treatment.", False),
        ("Members who use it each month", [0.85, 0.85, 0.9], S.PCT, "Not everyone comes every month. 85% is common.", False),
        ("Discount on products", [0.10, 0.10, 0.15], S.PCT, "Member perk on retail.", False),
        ("Average product spend a month", [30, 35, 45], S.USD, "Retail a typical member buys each month.", False),
        ("Discount on add-ons", [0.15, 0.15, 0.15], S.PCT, "Member perk on add-ons.", False),
        ("Average add-on spend a month", [15, 20, 25], S.USD, "Add-ons a typical member books each month.", False),
        ("Share of members on this tier", [0.5, 0.35, 0.15], S.PCT, "Your best guess. The three should add up to 100%.", True),
    ]
    for i, (lab, vals, fmt, n, key) in enumerate(tiers):
        r = 18 + i
        S.label(ws, "A%d" % r, lab)
        for col, v in zip("BCD", vals):
            c = S.inp(ws, "%s%d" % (col, r), v, fmt, key=key)
            c.alignment = Alignment(horizontal="right")
        S.note(ws, "E%d" % r, n)
    S.label(ws, "A29", "Mix check (should be 100%)")
    S.calc(ws, "B29", "=SUM(B28:D28)", S.PCT)
    S.calc(ws, "C29", '=IF(ROUND(B29,4)=1,"Good","Adjust the mix")')

    S.head(ws, 31, "What each member is worth", 5)
    worth = [
        ("Card fee", "={c}19*$B$14+$B$15", S.USD2, "On each monthly charge."),
        ("Product for the included service", "={c}22*{c}23", S.USD2, "Product cost times how often it's used."),
        ("Perks you give", "={c}25*{c}24+{c}27*{c}26", S.USD2, "The discounts, in dollars."),
        ("You keep per member, per month", "={c}19-{c}32-{c}33-{c}34", S.USD2, "Before rent and other monthly costs."),
        ("Member saves on each visit", "={c}20-{c}19", S.USD, "What makes the membership an easy yes."),
        ("Member's discount", "=IF({c}20=0,0,1-{c}19/{c}20)", S.PCT, "10% to 25% feels generous and stays profitable."),
        ("You keep per hour of treatment", "=IF({c}21*{c}23=0,0,{c}35/({c}21*{c}23/60))", S.USD, "Compare with what a regular hour earns you."),
    ]
    for i, (lab, f, fmt, n) in enumerate(worth):
        r = 32 + i
        S.label(ws, "A%d" % r, lab, bold=(r == 35))
        for col in "BCD":
            S.calc(ws, "%s%d" % (col, r), f.format(c=col), fmt, bold=(r == 35))
        S.note(ws, "E%d" % r, n)

    S.head(ws, 40, "The answers", 5)
    S.label(ws, "A41", "Average you keep per member (your mix)")
    S.calc(ws, "B41", "=IF(SUM(B28:D28)=0,0,SUMPRODUCT(B35:D35,B28:D28)/SUM(B28:D28))", S.USD2)
    S.label(ws, "A42", "Members to cover your monthly costs", bold=True)
    S.calc(ws, "B42", '=IF(B41<=0,"Raise prices",ROUNDUP(B11/B41,0))', S.NUM, big=True)
    S.note(ws, "E42", "Your break-even: the membership alone pays the bills.")
    S.label(ws, "A43", "Your monthly pay goal, before taxes")
    S.inp(ws, "B43", 4000, S.USD, key=True)
    S.note(ws, "E43", "What you want the membership to pay you each month.")
    S.label(ws, "A44", "Members to cover costs and your pay", bold=True)
    S.calc(ws, "B44", '=IF(B41<=0,"Raise prices",ROUNDUP((B11+B43)/B41,0))', S.NUM, big=True)
    S.label(ws, "A45", "Monthly dues from that many members")
    S.calc(ws, "B45", "=IF(ISNUMBER(B44),B44*SUMPRODUCT(B19:D19,B28:D28)/SUM(B28:D28),0)", S.USD)
    S.label(ws, "A46", "Treatment hours a month they need")
    S.calc(ws, "B46", "=IF(ISNUMBER(B44),B44*SUMPRODUCT(B21:D21,B23:D23,B28:D28)/SUM(B28:D28)/60,0)", S.NUM1)
    S.label(ws, "A47", "Hours you treat clients a month")
    S.inp(ws, "B47", 140, S.NUM)
    S.note(ws, "E47", "About 32 hours a week. Change it to your schedule.")
    S.label(ws, "A48", "Share of your hours members fill", bold=True)
    S.calc(ws, "B48", "=IF(B47=0,0,B46/B47)", S.PCT, bold=True)
    S.calc(ws, "C48", '=IF(B48>0.8,"Too full: raise prices or cap it",IF(B48<0.6,"Room to grow","A healthy balance"))')
    S.note(ws, "E48", "Under 60% leaves room for new clients. Over 80%, raise prices or cap the membership.")
    ws.freeze_panes = "B4"


def growth(ws):
    S.title(ws, "Your First Year", "How the membership grows with steady sign-ups and a few cancellations.")
    for col, w in zip("ABCDEFGHIJK", (12, 13, 13, 13, 13, 14, 16, 14, 14, 15, 14)):
        ws.column_dimensions[col].width = w
    S.head(ws, 4, "Assumptions", 11)
    S.label(ws, "A5", "Starting members")
    S.inp(ws, "C5", 0, S.NUM)
    S.label(ws, "A6", "Cancel each month")
    S.inp(ws, "C6", 0.04, S.PCT, key=True)
    S.note(ws, "E6", "3% to 5% a month is typical for a well-run membership.")
    S.label(ws, "A7", "Average dues")
    c = S.calc(ws, "C7", "=IF(SUM(Calculator!B28:D28)=0,0,SUMPRODUCT(Calculator!B19:D19,Calculator!B28:D28)/SUM(Calculator!B28:D28))", S.USD2)
    c.font = S.LINK
    S.label(ws, "A8", "You keep per member")
    S.calc(ws, "C8", "=Calculator!B41", S.USD2).font = S.LINK
    S.label(ws, "A9", "Monthly costs")
    S.calc(ws, "C9", "=Calculator!B11", S.USD).font = S.LINK
    S.note(ws, "E7", "Green numbers come from the Calculator tab.")

    heads = ["Month", "Start", "New sign-ups", "Cancelled", "End", "Dues", "You keep", "Monthly costs", "Net",
             "Running total", "Covers costs"]
    for i, h in enumerate(heads):
        c = ws.cell(row=11, column=i + 1, value=h.upper())
        c.font = S.HEAD
        c.border = S.LINE
        c.alignment = Alignment(horizontal="right" if i else "left", wrap_text=True)
    ws.row_dimensions[11].height = 28
    new = [4, 5, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6]
    for m in range(12):
        r = 12 + m
        ws.cell(row=r, column=1, value=m + 1).font = S.LABEL
        S.calc(ws, "B%d" % r, "=C5" if m == 0 else "=E%d" % (r - 1), S.NUM1)
        S.inp(ws, "C%d" % r, new[m], S.NUM)
        S.calc(ws, "D%d" % r, "=B%d*$C$6" % r, S.NUM1)
        S.calc(ws, "E%d" % r, "=B%d+C%d-D%d" % (r, r, r), S.NUM1)
        S.calc(ws, "F%d" % r, "=E%d*$C$7" % r, S.USD)
        S.calc(ws, "G%d" % r, "=E%d*$C$8" % r, S.USD)
        S.calc(ws, "H%d" % r, "=$C$9", S.USD)
        S.calc(ws, "I%d" % r, "=G%d-H%d" % (r, r), S.USD)
        S.calc(ws, "J%d" % r, "=I%d" % r if m == 0 else "=J%d+I%d" % (r - 1, r), S.USD)
        S.calc(ws, "K%d" % r, '=IF(I%d>=0,"Yes","Not yet")' % r)
        for col in range(1, 12):
            ws.cell(row=r, column=col).border = S.LINE
    S.note(ws, "C24", "New sign-ups are blue: change any month.")
    S.head(ws, 26, "After twelve months", 11)
    S.label(ws, "A27", "Members")
    S.calc(ws, "D27", "=E23", S.NUM1, bold=True)
    S.label(ws, "A28", "Monthly dues")
    S.calc(ws, "D28", "=F23", S.USD, bold=True)
    S.label(ws, "A29", "First month the membership covers your costs")
    S.calc(ws, "D29", '=IF(COUNTIF(K12:K23,"Yes")=0,"Not yet",INDEX(A12:A23,MATCH("Yes",K12:K23,0)))', S.NUM, bold=True)
    ws.freeze_panes = "A12"


def price_test(ws):
    S.title(ws, "Price Test", "Monthly profit from Tier 1 at different prices and member counts, after your monthly costs.")
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 16
    for col in "CDEFGH":
        ws.column_dimensions[col].width = 13
    ws["B4"] = "PRICE  ↓   MEMBERS  →"
    ws["B4"].font = S.HEAD
    members = [10, 20, 30, 40, 50, 60]
    for j, m in enumerate(members):
        S.inp(ws, "%s4" % "CDEFGH"[j], m, S.NUM).alignment = Alignment(horizontal="right")
    prices = list(range(89, 200, 10))
    for i, p in enumerate(prices):
        r = 5 + i
        S.inp(ws, "B%d" % r, p, S.USD)
        for j in range(len(members)):
            col = "CDEFGH"[j]
            S.calc(ws, "%s%d" % (col, r),
                   "={c}$4*($B{r}-($B{r}*Calculator!$B$14+Calculator!$B$15)-Calculator!$B$33-Calculator!$B$34)"
                   "-Calculator!$B$11".format(c=col, r=r), S.USD)
    last = 4 + len(prices)
    rng = "C5:H%d" % last
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"],
                                                  fill=PatternFill("solid", fgColor="F8DCDC")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["0"],
                                                  fill=PatternFill("solid", fgColor="E3F1E3")))
    S.note(ws, "B%d" % (last + 2), "Uses Tier 1's product, perk and card-fee numbers from the Calculator tab. "
                                   "Red: the membership doesn't cover your monthly costs yet. Green: it does.")
    S.note(ws, "B%d" % (last + 3), "The prices and member counts in blue can be changed.")


def make(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Start Here"
    start_sheet(ws, NAME, [
        "Go to the Calculator tab. Put in your monthly costs (rent, insurance, software and so on).",
        "Set up to three tiers: the price, the facial that's included, what it costs you to do, and the perks.",
        "Read The Answers: what each member is worth, how many you need to cover your costs, and how many to pay you.",
        "Open Your First Year to see how fast you get there with steady sign-ups and a few cancellations.",
        "Use Price Test to compare prices before you announce them.",
    ], [("Good to know", "The example numbers are modeled on a small studio in Los Angeles. They are a "
                         "starting point, not advice. Every number in blue is yours to change."),
        ("Made by", "Evelyn Romero, licensed esthetician, Lumevina Aesthetics. Questions? Message us on Etsy.")])
    calculator(wb.create_sheet("Calculator"))
    growth(wb.create_sheet("Your First Year"))
    price_test(wb.create_sheet("Price Test"))
    for w in wb.worksheets:
        w.sheet_properties.tabColor = S.ROSE
    wb.active = 1
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return S.recalc(path)


START = {
    "inside": [
        ("Spreadsheet", "The calculator as an Excel file. Opens in Excel, Google Sheets (File, Import) and Numbers."),
        ("Four tabs", "Start Here · Calculator (your costs, tiers and the answers) · Your First Year (a 12-month "
                      "growth plan) · Price Test (profit at different prices and member counts)"),
    ],
    "tips": [
        "Open the Calculator tab and replace the blue example numbers with yours. Yellow cells matter most.",
        "Try a few prices and perks until what you keep per member feels right and the member still saves.",
        "Check how many members cover your costs, and how many hours of your month they'd fill.",
        "Use Your First Year to set a sign-up goal, and Price Test before you announce prices.",
    ],
    "extra": [("h", "Google Sheets"),
              ("p", "In Google Drive, choose **New, File upload**, pick the file, then **Open with Google Sheets**. "
                    "Everything calculates the same way.")],
    "legal": False,
    "fillable": False,
}


def build(c):
    fn = "Membership Pricing Calculator.xlsx"
    c.after.append(lambda: make(c.path(fn)))
    c.add(fn)
    c.start_here(formats=())
