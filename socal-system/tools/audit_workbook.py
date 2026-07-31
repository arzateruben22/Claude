"""Cross-check a SoCal Checklist Budget Tracker workbook.

Passes:
  1. Checklist "Break N Cylinders/Beams" rows -> matching Whittier log entries
     (set #, break date, qty, cure days, tank).
  2. Reverse: performed log breaks (PASS/FAIL) -> checklist row on that date.
  3. Internal log sanity: times in date columns, 1900 break dates,
     missing tanks, set-name date prefix vs actual cast date.
  4. Invoice lab-line reconciliation for a given month tab:
     tests == checklist specimens broken that day.

Usage:
    python audit_workbook.py <workbook.xlsx> [--invoice "105 Invoice July '26 (New)"]

Output: findings printed to stdout, tab-separated; one line per issue with
tab + Excel row references. Exit code 0 always (findings are data, not errors).
"""
import argparse
import datetime
import re
import sys
import warnings

warnings.filterwarnings("ignore")
from openpyxl import load_workbook  # noqa: E402

CL_COLS = ["A", "B", "C", "D", "Inspector", "Project", "EA", "TaskOrder", "Date",
           "Time", "JobDesc", "LogNo", "Group", "Budget", "Dispatch", "QC",
           "DispReq", "ConcLog", "RptClient", "RptTeams", "DailyLog",
           "Timesheet", "Vision", "Status"]


def alnum(s):
    return re.sub(r"[^A-Z0-9]", "", str(s or "").upper())


def clean_checklist_setno(s):
    s = str(s or "")
    s = re.sub(r"(?i)[/\s]*\(?\d[\d,.\-\s]*\)?\s*psi.*$", "", s)
    s = re.sub(r"(?i)@\s*\d+\s*-?\s*days?.*$", "", s)
    return alnum(s)


def set_match(cl_value, log_value):
    a, b = clean_checklist_setno(cl_value), alnum(log_value)
    if not a or not b:
        return False
    if a == b:
        return True
    lo, hi = (a, b) if len(a) < len(b) else (b, a)
    # prefix match, but never let S1 match S19
    return hi.startswith(lo) and not (lo[-1].isdigit() and hi[len(lo)].isdigit())


def iso(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    return str(v or "")[:10]


def load_rows(ws, n_cols):
    out = []
    for i, row in enumerate(ws.iter_rows(min_row=1, max_col=n_cols, values_only=True), 1):
        out.append((i, list(row)))
    return out


def parse_break_row(desc):
    m = re.match(r"(?i)\s*break\s+(\d+)\s+(cylinders?|beams?)\s*@\s*(\d+)[\s-]*days?", str(desc or ""))
    if not m:
        return None
    qty, kind, days = int(m.group(1)), m.group(2).lower().rstrip("s"), int(m.group(3))
    tm = re.search(r"(?i)tank\s*[:\-]?\s*([^/\n(]+?)(?:\s*/|\s*$|\s*\()", str(desc))
    tank = alnum(re.sub(r"(?i)[\s\-]*\d[\d,]*\s*psi.*$", "", tm.group(1))) if tm else ""
    return qty, kind, days, tank


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--invoice", default=None, help="invoice tab name to reconcile")
    args = ap.parse_args()

    wb = load_workbook(args.workbook, data_only=True, read_only=True)
    findings = []

    def add(sev, where, row, msg):
        findings.append((sev, where, row, msg))

    # ---- load checklist ----
    cl = []
    ws = wb["SoCal CList"]
    for i, row in enumerate(ws.iter_rows(min_row=5, max_col=24, values_only=True), 5):
        rec = dict(zip(CL_COLS, row))
        rec["row"], rec["date"] = i, iso(rec["Date"])
        cl.append(rec)

    # ---- load logs ----
    logs = {}
    for kind, tab, setcol in [("cylinder", "Whittier Cylinder Log", 14),
                              ("beam", "Whittier Beam Log", 15)]:
        rows = []
        ws = wb[tab]
        breakcol = 11 if kind == "cylinder" else 12
        for i, row in enumerate(ws.iter_rows(min_row=5, max_col=17, values_only=True), 5):
            if not row[0]:
                continue
            rows.append({
                "row": i, "breaklog": row[0], "tank": row[1],
                "castdate": row[3], "cast": iso(row[3]), "qty": row[9],
                "days": row[10], "breakdate": iso(row[breakcol]),
                "result": str(row[13] or ""), "set": row[setcol] or "",
            })
        logs[kind] = rows

    # ---- pass 3: internal sanity ----
    for kind, rows in logs.items():
        for r in rows:
            if r["castdate"] and not str(r["castdate"])[:2] == "20":
                add("High", f"{kind} log", r["row"],
                    f"Cast Date is not a date ({r['castdate']!r}) — set {r['set']}")
            if r["breakdate"].startswith("1900") or r["breakdate"].startswith("1899"):
                add("High", f"{kind} log", r["row"],
                    f"Break date computed as {r['breakdate']} — bad cast date? set {r['set']}")
            if str(r["tank"] or "").strip() in ("?", ""):
                if r["cast"] >= "2026-01-01":
                    add("Low", f"{kind} log", r["row"], f"Tank missing/'?' — set {r['set']}")
            m = re.match(r"(\d{2})(\d{2})(\d{2})-", str(r["set"]))
            if m and r["cast"].startswith("20"):
                mm, dd, yy = m.groups()
                want = f"20{yy}-{mm}-{dd}"
                got = r["cast"]
                if want != got:
                    try:
                        delta = abs((datetime.date.fromisoformat(want)
                                     - datetime.date.fromisoformat(got)).days)
                    except ValueError:
                        delta = 999
                    if delta > 1:
                        add("Med", f"{kind} log", r["row"],
                            f"Set name date prefix {mm}/{dd}/{yy} != cast date {got} — set {r['set']}")

    # ---- pass 1: checklist break rows -> logs ----
    for rec in cl:
        parsed = parse_break_row(rec["JobDesc"])
        if not parsed:
            continue
        qty, kind, days, tank = parsed
        pool = [r for r in logs[kind] if r["breakdate"] == rec["date"]]

        def score(r):
            s = 0
            if set_match(rec["LogNo"], r["set"]):
                s += 4
            if tank and alnum(r["tank"]) == tank:
                s += 2
            try:
                if int(float(r["qty"])) == qty:
                    s += 1
                if int(float(r["days"])) == days:
                    s += 1
            except (TypeError, ValueError):
                pass
            return s

        best = max(pool, key=score, default=None)
        if best is None or score(best) == 0:
            anywhere = [r for r in logs[kind] if set_match(rec["LogNo"], r["set"])]
            if anywhere:
                same_days = [r for r in anywhere
                             if str(r["days"]) and int(float(r["days"])) == days]
                tgt = same_days[0] if same_days else anywhere[0]
                add("Med", "SoCal CList", rec["row"],
                    f"Break date mismatch: checklist {rec['date']} vs log row "
                    f"{tgt['row']} breaking {tgt['breakdate']} ({tgt['set']})")
            else:
                add("High", "SoCal CList", rec["row"],
                    f"No {kind} log entry found for {str(rec['LogNo'])[:50]!r} on {rec['date']}")
            continue
        diffs = []
        if not set_match(rec["LogNo"], best["set"]):
            diffs.append(f"set#: CL {str(rec['LogNo'])[:40]!r} vs log {best['set']!r}")
        if tank and alnum(best["tank"]) and alnum(best["tank"]) != tank:
            diffs.append(f"tank: CL {tank} vs log {best['tank']}")
        try:
            if int(float(best["qty"])) != qty:
                diffs.append(f"qty: CL {qty} vs log {best['qty']}")
            if int(float(best["days"])) != days:
                diffs.append(f"days: CL {days} vs log {best['days']}")
        except (TypeError, ValueError):
            pass
        if diffs:
            add("Med", "SoCal CList", rec["row"],
                f"vs {best['breaklog']} (log row {best['row']}): " + "; ".join(diffs))

    # ---- pass 2: reverse ----
    break_rows_by_date = {}
    for rec in cl:
        parsed = parse_break_row(rec["JobDesc"])
        if parsed:
            break_rows_by_date.setdefault((parsed[1], rec["date"]), []).append(rec)
    today = datetime.date.today().isoformat()
    for kind, rows in logs.items():
        for r in rows:
            if not ("2026-01-01" <= r["breakdate"] <= today):
                continue
            if r["result"].upper() not in ("PASS", "FAIL", "PASSED", "FAILED"):
                continue
            cands = break_rows_by_date.get((kind, r["breakdate"]), [])
            if not any(set_match(c["LogNo"], r["set"]) for c in cands):
                add("Med", f"{kind} log", r["row"],
                    f"Performed break ({r['set']}, {r['breakdate']}, result "
                    f"{r['result']}) has no checklist row that day")

    # ---- pass 4: invoice reconciliation ----
    if args.invoice and args.invoice in wb.sheetnames:
        inv = {}
        ws = wb[args.invoice]
        for row in ws.iter_rows(min_row=13, max_row=260, max_col=11, values_only=True):
            insp, date, desc, tests = row[1], iso(row[3]), str(row[6] or ""), row[10]
            if insp != "Whittier Lab" or "Break" not in desc:
                continue
            if not isinstance(tests, (int, float)) or tests > 50:
                continue
            kind = "cylinder" if "Cylinder" in desc else "beam"
            inv[(kind, date)] = inv.get((kind, date), 0) + (tests or 0)
        cl_daily = {}
        for rec in cl:
            parsed = parse_break_row(rec["JobDesc"])
            if parsed and rec["Inspector"] == "Whittier Lab" and "105" in str(rec["Project"]):
                cl_daily[(parsed[1], rec["date"])] = \
                    cl_daily.get((parsed[1], rec["date"]), 0) + parsed[0]
        for key in sorted(set(inv) | set(cl_daily)):
            a, b = cl_daily.get(key, 0), inv.get(key, 0)
            if a != b and b != 0:
                add("Med", args.invoice, 0,
                    f"{key[1]} {key[0]}s: checklist {a} specimens vs invoice {b} tests")

    # ---- report ----
    order = {"High": 0, "Med": 1, "Low": 2}
    findings.sort(key=lambda f: (order.get(f[0], 3), f[1], f[2]))
    print(f"{len(findings)} findings")
    for sev, where, row, msg in findings:
        print(f"{sev}\t{where}\trow {row}\t{msg}")


if __name__ == "__main__":
    main()
